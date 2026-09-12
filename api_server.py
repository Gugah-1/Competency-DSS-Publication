from __future__ import annotations

import io
import json
import shutil
import sqlite3
from datetime import date, timedelta, datetime
from contextvars import ContextVar
import secrets
import time
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Query, Request
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from config import BUILD, DB_BACKEND, SESSION_COOKIE_NAME, SESSION_COOKIE_SECURE, ROLE_LABELS
from database import backend_info, production_preflight
from assessment_engine import ENGINE_VERSION, assessment_readiness, run_employee_assessment
from auth_service import (
    authenticate, change_password, create_session, hash_password, password_policy,
    revoke_session, revoke_user_sessions, role_permissions, session_user, set_temporary_password,
)
from backup_service import (
    BACKUP_DIR, create_backup, ensure_daily_backup, list_backups, prune_backups, restore_backup, validate_backup,
)
from employee_purge import employee_purge_impact as calculate_employee_purge_impact, is_test_employee_code, purge_employee_records
from reporting_service import (
    analytics_frames, list_departments, management_ai_summary, management_analytics, refresh_monthly_snapshot, trend_analytics,
)
from workflow_v226 import (
    ACTION_PRIORITIES, ACTION_STATUSES, ACTION_TYPES, ALLOWED_EVIDENCE_EXT, MAX_EVIDENCE_BYTES,
    action_completion_ready, assert_entity, create_action as workflow_create_action,
    infer_entity_context, refresh_notifications, update_action as workflow_update_action, _safe_filename,
)
from migration_v22_10_3_to_v22_10_4 import migrate as migrate_v22_10_4
from migration_v22_10_4_to_v22_10_5 import migrate as migrate_v22_10_5

ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "data" / "competency_dss.db"
WEB_DIR = ROOT / "web"
TEMPLATE_DIR = ROOT / "templates"

# Idempotent startup migration. Existing action/evidence rows are preserved.
migrate_v22_10_4(DB_PATH)
migrate_v22_10_5(DB_PATH)


app = FastAPI(
    title="Competency DSS API",
    version="22.10.5",
    description="Authenticated LAN backend with management analytics and reporting for the Data-Driven Decision Support Framework."
)

if WEB_DIR.exists():
    app.mount("/app/static", StaticFiles(directory=WEB_DIR), name="static")

@app.get("/", include_in_schema=False)
def web_home():
    index_path = WEB_DIR / "index.html"
    if not index_path.exists():
        raise HTTPException(status_code=404, detail="Web frontend not found")
    return FileResponse(index_path)


def query_df(sql: str, params: tuple = ()) -> pd.DataFrame:
    with sqlite3.connect(DB_PATH) as con:
        return pd.read_sql_query(sql, con, params=params)


def today() -> date:
    return date.today()

class AIQuery(BaseModel):
    question: str
    conversation_id: str | None = None
    reset_context: bool = False


def is_management_analytics_question(question: str) -> bool:
    ql = (question or "").strip().lower()
    markers = (
        "management summary", "management analytics", "ringkasan management", "ringkasan manajemen",
        "ringkas kondisi", "bulan ini", "monthly report", "laporan bulanan", "90 hari ke depan",
        "requirement coverage rate", "training completion rate", "action closure rate",
        "potential gap resolution", "department comparison", "perbandingan department", "perbandingan departemen"
    )
    return any(marker in ql for marker in markers)


class EmployeeCreate(BaseModel):
    employee_code: str
    employee_name: str
    position_id: int
    department: str | None = None
    join_date: date | None = None
    employment_status: str = "Active"
    remarks: str | None = None


class EmployeeUpdate(BaseModel):
    employee_name: str
    department: str | None = None
    join_date: date | None = None
    employment_status: str = "Active"
    exit_date: date | None = None
    status_reason: str | None = None
    remarks: str | None = None


class PositionChange(BaseModel):
    position_id: int
    effective_date: date
    change_reason: str = "Position Change"
    remarks: str | None = None


class EmployeeDeactivate(BaseModel):
    effective_date: date
    employment_status: str = "Resigned"
    reason: str = "Resignation"
    remarks: str | None = None


class EmployeePurge(BaseModel):
    confirmation_code: str
    reason: str


class TrainingCatalogCreate(BaseModel):
    training_name: str
    category: str | None = None


class TrainingCatalogUpdate(BaseModel):
    training_name: str
    category: str | None = None
    active_flag: int = 1


class TrainingHistoryCreate(BaseModel):
    employee_pk: int
    training_id: int
    training_date: date | None = None
    completion_date: date | None = None
    result_status: str = "Completed"
    provider: str | None = None
    valid_until: date | None = None
    certificate_reference: str | None = None
    evidence_note: str | None = None
    remarks: str | None = None


class TrainingHistoryUpdate(TrainingHistoryCreate):
    pass


class TrainingRequirementCreate(BaseModel):
    position_id: int
    training_id: int
    requirement_type: str = "Mandatory"
    requirement_code: str | None = None
    delivery_type: str | None = None
    regulatory_flag: str | None = None
    requirement_source: str = "Manual TCD"
    effective_from: date | None = None
    effective_to: date | None = None
    remarks: str | None = None


class TrainingRequirementUpdate(TrainingRequirementCreate):
    requirement_status: str = "Active"


class RequirementDeactivate(BaseModel):
    effective_to: date
    reason_code: str = "Requirement Changed"
    remarks: str | None = None


class CertificationCreate(BaseModel):
    employee_pk: int
    certification_name: str
    certificate_number: str | None = None
    issuer: str | None = None
    issue_date: date | None = None
    expiry_date: date | None = None
    certification_status: str | None = None
    remarks: str | None = None


class CertificationUpdate(CertificationCreate):
    renewal_status: str | None = None
    reason_code: str | None = None


class CertificationRenew(BaseModel):
    certificate_number: str | None = None
    issuer: str | None = None
    issue_date: date
    expiry_date: date | None = None
    remarks: str | None = None


class CertificationNotRenewed(BaseModel):
    reason_code: str
    remarks: str | None = None


class AssessmentConfirm(BaseModel):
    remarks: str | None = None


class MappingValidationCreate(BaseModel):
    requirement_id: int
    employee_pk: int | None = None
    scope: str = "Employee"
    decision: str
    notes: str | None = None
    evidence_reference: str | None = None


class AssessmentOverrideCreate(BaseModel):
    employee_pk: int
    requirement_id: int
    override_status: str
    reason: str
    evidence_reference: str | None = None


class RevokeDecision(BaseModel):
    remarks: str | None = None


class ImportConfirm(BaseModel):
    accept_warnings: bool = True
    duplicate_action: str = "skip"


class ActionCreate(BaseModel):
    employee_pk: int
    requirement_id: int | None = None
    action_type: str
    title: str | None = None
    priority: str = "Medium"
    recommendation: str | None = None
    pic_user_id: str | None = None
    due_date: date | None = None
    remarks: str | None = None
    source_decision_id: int | None = None


class ActionUpdate(BaseModel):
    status: str | None = None
    pic_user_id: str | None = None
    due_date: date | None = None
    priority: str | None = None
    remarks: str | None = None
    no_action_reason: str | None = None


class ActionDelete(BaseModel):
    reason: str


class ActionRestore(BaseModel):
    reason: str


class NotificationRead(BaseModel):
    read: bool = True


class LoginRequest(BaseModel):
    username: str
    password: str


class PasswordChangeRequest(BaseModel):
    current_password: str
    new_password: str


class UserCreateRequest(BaseModel):
    username: str
    display_name: str
    role: str
    temporary_password: str | None = None


class UserResetPasswordRequest(BaseModel):
    temporary_password: str | None = None


class RestoreBackupRequest(BaseModel):
    confirm: str



_AUTH_USER: ContextVar[dict[str, Any] | None] = ContextVar("dss_auth_user", default=None)
_AUTH_TOKEN: ContextVar[str | None] = ContextVar("dss_auth_token", default=None)
_LOGIN_ATTEMPTS: dict[str, list[float]] = {}
_AI_CONTEXTS: dict[str, dict[str, Any]] = {}


def current_user() -> dict[str, Any]:
    user = _AUTH_USER.get()
    if not user:
        raise HTTPException(status_code=401, detail="Authentication required.")
    return dict(user)


def required_permission(method: str, path: str) -> str | None:
    if path.startswith("/api/auth/"):
        return None
    if method == "GET":
        if path.startswith("/api/system/backups"):
            return "backup_manage"
        if path.startswith("/api/system/users"):
            return "user_manage"
        if path.startswith("/api/audit"):
            return "audit_view"
        if "/export" in path or path.startswith("/api/export/"):
            return "export"
        return "read"
    if path.startswith("/api/system/backups"):
        return "backup_manage"
    if path.startswith("/api/system/users"):
        return "user_manage"
    if path.startswith("/api/employees/") and "/assessment/" in path:
        return "assessment_manage"
    if path.startswith("/api/employees"):
        return "employee_manage"
    if path.startswith("/api/training-history"):
        return "training_history_manage"
    if path.startswith("/api/certifications"):
        return "certification_manage"
    if path.startswith("/api/training-requirements") or path.startswith("/api/training-catalog"):
        return "requirement_manage"
    if path.startswith("/api/assessment-validations"):
        return "validation_manage"
    if path.startswith("/api/assessment-overrides"):
        return "override_manage"
    if path.startswith("/api/actions"):
        return "action_manage"
    if path.startswith("/api/evidence"):
        return "evidence_manage"
    if path.startswith("/api/data-management/import"):
        return "import_manage"
    if path.startswith("/api/refresh"):
        return "refresh"
    if path.startswith("/api/notifications") or path.startswith("/api/ai/query"):
        return "read"
    if path.startswith("/api/ai/export"):
        return "export"
    return "system_write"


def _public_path(path: str) -> bool:
    return path == "/" or path == "/health" or path == "/api/build" or path.startswith("/app/static/") or path == "/api/auth/login"


@app.middleware("http")
async def authentication_middleware(request: Request, call_next):
    path = request.url.path
    if _public_path(path):
        return await call_next(request)
    token = request.cookies.get(SESSION_COOKIE_NAME)
    with sqlite3.connect(DB_PATH) as con:
        user = session_user(con, token, touch=True)
        con.commit()
    if not user:
        return JSONResponse(status_code=401, content={"detail": "Session tidak tersedia atau sudah berakhir. Silakan login kembali."})
    # A temporary password must be changed before operational data can be opened.
    if user.get("must_change_password") and path not in {"/api/auth/me", "/api/auth/change-password", "/api/auth/logout"}:
        return JSONResponse(status_code=403, content={"detail": "PASSWORD_CHANGE_REQUIRED"})
    permission = required_permission(request.method.upper(), path)
    if permission and permission not in set(user.get("permissions") or []):
        return JSONResponse(status_code=403, content={"detail": f"Permission denied: {permission}"})
    uctx = _AUTH_USER.set(user)
    tctx = _AUTH_TOKEN.set(token)
    try:
        return await call_next(request)
    finally:
        _AUTH_USER.reset(uctx)
        _AUTH_TOKEN.reset(tctx)


def write_audit(user: dict[str, Any], action: str, dataset_type: str = "", target: str = "", details: str = "") -> None:
    with sqlite3.connect(DB_PATH) as con:
        con.execute(
            "INSERT INTO audit_log(user_id, action, dataset_type, target, details) VALUES (?,?,?,?,?)",
            (user["user_id"], action, dataset_type or None, target or None, details or None),
        )
        con.commit()


def iso(v: date | None) -> str | None:
    return v.isoformat() if v else None


def require_row(con: sqlite3.Connection, table: str, key: str, value: Any, label: str) -> sqlite3.Row:
    con.row_factory = sqlite3.Row
    row = con.execute(f"SELECT * FROM {table} WHERE {key}=?", (value,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail=f"{label} tidak ditemukan.")
    return row


def mark_employee_for_reassessment(con: sqlite3.Connection, employee_pk: int, user_id: str) -> None:
    row = con.execute("SELECT assessment_state FROM employees WHERE employee_pk=?", (employee_pk,)).fetchone()
    if not row:
        return
    state = row[0] or "Assessed"
    # Before the first assessment, added evidence is part of initial data completion.
    # Only an already Assessed employee transitions into Pending Reassessment.
    if state == "Assessed":
        con.execute(
            "UPDATE employees SET assessment_state='Pending Reassessment', updated_by=?, updated_at=CURRENT_TIMESTAMP WHERE employee_pk=?",
            (user_id, employee_pk),
        )


def mark_position_for_reassessment(con: sqlite3.Connection, position_id: int, user_id: str) -> int:
    cur = con.execute(
        """UPDATE employees SET assessment_state='Pending Reassessment', updated_by=?, updated_at=CURRENT_TIMESTAMP
           WHERE active_flag=1 AND current_position_id=? AND COALESCE(assessment_state,'Assessed')='Assessed'""",
        (user_id, position_id),
    )
    return int(cur.rowcount or 0)


def certification_status_for(expiry_date: date | None, requested: str | None = None) -> str:
    if requested and requested not in {"", "Auto"}:
        return requested
    if not expiry_date:
        return "Active"
    d = today()
    if expiry_date < d:
        return "Expired"
    if expiry_date <= d + timedelta(days=90):
        return "Near Expiry"
    return "Active"


def renewal_status_for(validity_status: str | None, requested: str | None = None) -> str:
    """Normalize renewal workflow without treating expiry as an automatic renewal decision."""
    raw = str(requested or '').strip()
    aliases = {
        'Planned': 'Renewal Planned',
    }
    raw = aliases.get(raw, raw)
    explicit = {
        'Renewal Planned', 'Renewal Required', 'Renewal Pending',
        'Completed', 'Not Renewed', 'No Action Required',
        'Renewal Review Required',
    }
    if raw in explicit:
        return raw
    if str(validity_status or '') in {'Near Expiry', 'Expired'}:
        # Near-expiry/expired means human review is due, not that renewal is mandatory.
        return 'Renewal Review Required'
    return 'Not Due'


def resolve_certification_catalog(con: sqlite3.Connection, name: str, issuer: str | None = None) -> int:
    clean = name.strip()
    if not clean:
        raise HTTPException(status_code=400, detail="Nama sertifikasi wajib diisi.")
    row = con.execute("SELECT certification_catalog_id FROM certification_catalog WHERE lower(certification_name)=lower(?)", (clean,)).fetchone()
    if row:
        return int(row[0])
    cur = con.execute(
        "INSERT INTO certification_catalog(certification_name, issuer_type, active_flag) VALUES (?,?,1)",
        (clean, issuer),
    )
    return int(cur.lastrowid)


@app.post("/api/auth/login")
def login(payload: LoginRequest, request: Request):
    key = request.client.host if request.client else "unknown"
    now = time.time()
    attempts = [t for t in _LOGIN_ATTEMPTS.get(key, []) if now - t < 300]
    if len(attempts) >= 5:
        raise HTTPException(status_code=429, detail="Terlalu banyak login gagal. Coba lagi dalam beberapa menit.")
    with sqlite3.connect(DB_PATH) as con:
        user = authenticate(con, payload.username, payload.password)
        if not user:
            attempts.append(now); _LOGIN_ATTEMPTS[key] = attempts
            con.execute("INSERT INTO audit_log(user_id,action,dataset_type,target,details) VALUES (NULL,'LOGIN_FAILED','Security',?,?)", (payload.username[:100], f"ip={key}"))
            con.commit()
            raise HTTPException(status_code=401, detail="Username atau password tidak sesuai.")
        _LOGIN_ATTEMPTS.pop(key, None)
        token = create_session(con, user["user_id"], key, request.headers.get("user-agent"))
        con.execute("INSERT INTO audit_log(user_id,action,dataset_type,target,details) VALUES (?,'LOGIN_SUCCESS','Security',?,?)", (user["user_id"], user["username"], f"ip={key}"))
        con.commit()
    response = JSONResponse({"status":"success","user":user,"role_label":ROLE_LABELS.get(user['role'],user['role'])})
    response.set_cookie(SESSION_COOKIE_NAME, token, httponly=True, samesite="lax", secure=SESSION_COOKIE_SECURE, max_age=12*3600, path="/")
    return response


@app.get("/api/auth/me")
def auth_me() -> dict[str, Any]:
    user = current_user()
    return {"user": user, "role_label": ROLE_LABELS.get(user["role"], user["role"])}


@app.post("/api/auth/logout")
def logout():
    user = current_user(); token = _AUTH_TOKEN.get()
    with sqlite3.connect(DB_PATH) as con:
        revoke_session(con, token); con.commit()
    write_audit(user, "LOGOUT", "Security", user.get("username") or "", "")
    response = JSONResponse({"status":"success"})
    response.delete_cookie(SESSION_COOKIE_NAME, path="/")
    return response


@app.post("/api/auth/change-password")
def auth_change_password(payload: PasswordChangeRequest):
    user = current_user()
    try:
        with sqlite3.connect(DB_PATH) as con:
            change_password(con, user["user_id"], payload.current_password, payload.new_password)
            con.execute("INSERT INTO audit_log(user_id,action,dataset_type,target,details) VALUES (?,'CHANGE_PASSWORD','Security',?,?)", (user["user_id"], user.get("username"), "all sessions revoked"))
            con.commit()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    response = JSONResponse({"status":"success","relogin_required":True})
    response.delete_cookie(SESSION_COOKIE_NAME, path="/")
    return response


@app.get("/api/system/users")
def list_system_users() -> list[dict[str, Any]]:
    df=query_df("SELECT user_id,username,display_name,role,active_flag,must_change_password,last_login_at,password_changed_at,created_at,updated_at FROM app_users ORDER BY active_flag DESC,display_name")
    rows=df.fillna("").to_dict(orient="records")
    for r in rows: r["role_label"]=ROLE_LABELS.get(r.get("role"),r.get("role"))
    return rows


@app.post("/api/system/users")
def create_system_user(payload: UserCreateRequest) -> dict[str, Any]:
    user=current_user()
    role=payload.role.strip()
    if role not in ROLE_LABELS: raise HTTPException(status_code=400,detail="Role hanya Supervisor TCD atau HRD.")
    username=payload.username.strip().lower()
    if not username: raise HTTPException(status_code=400,detail="Username wajib diisi.")
    temp=payload.temporary_password or (secrets.token_urlsafe(12)+"A1!")
    ok,msg=password_policy(temp)
    if not ok: raise HTTPException(status_code=400,detail=msg)
    uid="U"+datetime.now().strftime("%y%m%d%H%M%S")+secrets.token_hex(2).upper()
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.execute("INSERT INTO app_users(user_id,display_name,role,active_flag,username,password_hash,must_change_password) VALUES (?,?,?,?,?,?,1)", (uid,payload.display_name.strip(),role,1,username,hash_password(temp)))
            con.commit()
    except sqlite3.IntegrityError:
        raise HTTPException(status_code=409,detail="Username sudah digunakan.")
    write_audit(user,"CREATE_USER","Security",uid,f"username={username}; role={role}")
    return {"status":"success","user_id":uid,"username":username,"temporary_password":temp,"must_change_password":True}


@app.post("/api/system/users/{user_id}/reset-password")
def reset_system_user_password(user_id: str, payload: UserResetPasswordRequest) -> dict[str, Any]:
    actor=current_user(); temp=payload.temporary_password or (secrets.token_urlsafe(12)+"A1!")
    try:
        with sqlite3.connect(DB_PATH) as con:
            if not con.execute("SELECT 1 FROM app_users WHERE user_id=?",(user_id,)).fetchone(): raise HTTPException(status_code=404,detail="User tidak ditemukan.")
            set_temporary_password(con,user_id,temp); con.commit()
    except ValueError as exc: raise HTTPException(status_code=400,detail=str(exc))
    write_audit(actor,"RESET_USER_PASSWORD","Security",user_id,"temporary password issued; sessions revoked")
    return {"status":"success","user_id":user_id,"temporary_password":temp,"must_change_password":True}


@app.post("/api/system/users/{user_id}/deactivate")
def deactivate_system_user(user_id: str) -> dict[str, Any]:
    actor=current_user()
    if user_id==actor["user_id"]: raise HTTPException(status_code=400,detail="Tidak dapat menonaktifkan akun yang sedang digunakan.")
    with sqlite3.connect(DB_PATH) as con:
        if not con.execute("SELECT 1 FROM app_users WHERE user_id=?",(user_id,)).fetchone(): raise HTTPException(status_code=404,detail="User tidak ditemukan.")
        con.execute("UPDATE app_users SET active_flag=0,updated_at=CURRENT_TIMESTAMP WHERE user_id=?",(user_id,)); revoke_user_sessions(con,user_id); con.commit()
    write_audit(actor,"DEACTIVATE_USER","Security",user_id,"")
    return {"status":"success","user_id":user_id}


@app.get("/api/system/backups")
def system_backups(limit: int=50) -> list[dict[str, Any]]:
    return list_backups(limit)


@app.post("/api/system/backups")
def system_create_backup() -> dict[str, Any]:
    user=current_user(); result=create_backup("manual",user["user_id"]); prune_backups(30)
    write_audit(user,"CREATE_BACKUP","System Backup",result["filename"],f"size={result['size']}")
    return result


@app.get("/api/system/backups/{filename}/validate")
def system_validate_backup(filename: str) -> dict[str, Any]:
    try: return validate_backup(filename)
    except (ValueError,FileNotFoundError) as exc: raise HTTPException(status_code=400,detail=str(exc))


@app.get("/api/system/backups/{filename}/download")
def system_download_backup(filename: str):
    p=(BACKUP_DIR/filename).resolve()
    if BACKUP_DIR.resolve() not in p.parents or not p.exists(): raise HTTPException(status_code=404,detail="Backup tidak ditemukan.")
    return FileResponse(p,filename=p.name,media_type="application/zip")


@app.post("/api/system/backups/{filename}/restore")
def system_restore_backup(filename: str, payload: RestoreBackupRequest):
    user=current_user()
    if payload.confirm != "RESTORE": raise HTTPException(status_code=400,detail="Ketik RESTORE untuk mengonfirmasi operasi restore.")
    try: result=restore_backup(filename,user["user_id"])
    except (ValueError,FileNotFoundError,RuntimeError) as exc: raise HTTPException(status_code=400,detail=str(exc))
    # Restore revokes sessions; this audit is best-effort because the restored DB may not contain this current user session.
    try: write_audit(user,"RESTORE_BACKUP","System Backup",filename,f"safety={result['safety_backup']}")
    except Exception: pass
    response=JSONResponse(result); response.delete_cookie(SESSION_COOKIE_NAME,path="/"); return response


@app.on_event("startup")
def v229_startup_foundation():
    try:
        ensure_daily_backup("system-startup"); prune_backups(30)
    except Exception as exc:
        print("[v22.10.5] Daily backup warning:",exc)
    try:
        with sqlite3.connect(DB_PATH) as con:
            refresh_monthly_snapshot(con, created_by="system-startup")
            con.commit()
    except Exception as exc:
        print("[v22.10.5] Reporting snapshot warning:",exc)


@app.get("/api/build")
def build_info() -> dict[str, Any]:
    return {"build": BUILD, "frontend": "custom-web", "backend": "fastapi", "database": str(DB_PATH.name), "db_backend": DB_BACKEND, "database_info": backend_info()}

@app.get("/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "database": DB_PATH.exists(), "db_backend": DB_BACKEND, "database_info": backend_info()}


@app.get("/api/system/database")
def database_status() -> dict[str, Any]:
    return production_preflight()

@app.get("/api/audit")
def audit_log(limit: int = 100) -> list[dict[str, Any]]:
    df = query_df(
        """SELECT a.audit_id, a.created_at, a.user_id, COALESCE(u.display_name, 'Shared LAN User') AS display_name,
                  COALESCE(u.role, 'shared_user') AS role, a.action, a.dataset_type, a.target, a.details
           FROM audit_log a LEFT JOIN app_users u ON u.user_id=a.user_id
           ORDER BY a.created_at DESC, a.audit_id DESC LIMIT ?""",
        (min(max(limit, 1), 500),),
    ).fillna("")
    return df.to_dict(orient="records")

@app.post("/api/refresh")
def refresh_data() -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        counts = {
            "employees": int(con.execute("SELECT COUNT(*) FROM employees WHERE active_flag=1").fetchone()[0]),
            "certifications": int(con.execute("SELECT COUNT(*) FROM certifications").fetchone()[0]),
            "training_history": int(con.execute("SELECT COUNT(*) FROM training_history").fetchone()[0]),
            "training_requirements": int(con.execute("SELECT COUNT(*) FROM training_requirements").fetchone()[0]),
        }
    user = current_user()
    write_audit(user, "REFRESH_DATA", details=str(counts))
    return {"status": "ok", "refreshed": True, "counts": counts, "user": user}

@app.get("/api/reference/positions")
def reference_positions() -> list[dict[str, Any]]:
    return query_df("SELECT position_id, position_name, COALESCE(department,'') department FROM positions WHERE active_flag=1 ORDER BY position_name").fillna("").to_dict(orient="records")


@app.get("/api/reference/training-catalog")
def reference_training_catalog() -> list[dict[str, Any]]:
    return query_df("SELECT training_id, training_name, COALESCE(category,'') category FROM training_catalog WHERE active_flag=1 ORDER BY training_name").fillna("").to_dict(orient="records")


@app.get("/api/reference/certification-catalog")
def reference_certification_catalog() -> list[dict[str, Any]]:
    return query_df("SELECT certification_catalog_id, certification_name, COALESCE(issuer_type,'') issuer_type FROM certification_catalog WHERE active_flag=1 ORDER BY certification_name").fillna("").to_dict(orient="records")


@app.get("/api/employees")
def employees(limit: int = 1000, include_inactive: bool = False) -> list[dict[str, Any]]:
    limit = max(1, min(limit, 5000))
    where = "" if include_inactive else "WHERE e.active_flag = 1"
    df = query_df(
        f"""
        SELECT e.employee_pk, e.employee_code, e.employee_name, e.current_position_id,
               p.position_name, COALESCE(e.department,p.department,'') AS department,
               e.join_date, COALESCE(e.employment_status, CASE WHEN e.active_flag=1 THEN 'Active' ELSE 'Inactive' END) employment_status,
               e.exit_date, e.status_reason, e.remarks, e.assessment_state, e.initial_data_confirmed,
               e.initial_data_confirmed_at, e.initial_data_confirmed_by, e.last_assessment_at,
               e.last_assessment_run_id, e.assessment_version, e.record_status, e.active_flag, e.updated_at,
               (SELECT COUNT(*) FROM training_requirements tr WHERE tr.active_flag=1 AND COALESCE(tr.requirement_status,'Active')='Active' AND COALESCE(tr.position_standard_id,tr.position_id)=e.current_position_id) AS requirement_count,
               (SELECT COUNT(*) FROM certifications c WHERE c.employee_pk=e.employee_pk AND COALESCE(c.record_status,'Active')<>'Archived') AS certification_count,
               (SELECT COUNT(*) FROM actions a WHERE a.employee_pk=e.employee_pk AND COALESCE(a.is_deleted,0)=0 AND a.status NOT IN ('Completed','Closed','Cancelled','No Action Required')) AS open_action_count,
               (SELECT COUNT(*) FROM decision_assessment_current da WHERE da.employee_pk=e.employee_pk AND da.final_status='Potential Gap') AS potential_gap_count,
               (SELECT COUNT(*) FROM decision_assessment_current da WHERE da.employee_pk=e.employee_pk AND da.final_status='Validation Required') AS validation_count,
               (SELECT COUNT(*) FROM decision_assessment_current da WHERE da.employee_pk=e.employee_pk AND da.final_status='Covered') AS covered_count
        FROM employees e
        LEFT JOIN positions p ON p.position_id = e.current_position_id
        {where}
        ORDER BY e.active_flag DESC, e.employee_name
        LIMIT ?
        """,
        (limit,),
    )
    return df.fillna("").to_dict(orient="records")


@app.post("/api/employees")
def create_employee(payload: EmployeeCreate) -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        con.execute("PRAGMA foreign_keys=ON")
        pos = require_row(con, "positions", "position_id", payload.position_id, "Position")
        code = payload.employee_code.strip()
        name = payload.employee_name.strip()
        if not code or not name:
            raise HTTPException(status_code=400, detail="Employee ID dan nama wajib diisi.")
        if con.execute("SELECT 1 FROM employees WHERE lower(employee_code)=lower(?)", (code,)).fetchone():
            raise HTTPException(status_code=409, detail="Employee ID sudah digunakan.")
        department = (payload.department or pos["department"] or "").strip() or None
        cur = con.execute(
            """INSERT INTO employees(employee_code, employee_name, current_position_id, active_flag, department, join_date,
                   employment_status, remarks, assessment_state, initial_data_confirmed, record_status, created_by, updated_by)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (code, name, payload.position_id, 1, department, iso(payload.join_date), payload.employment_status or "Active",
             payload.remarks, "Pending Initial Assessment", 0, "Active", user["user_id"], user["user_id"]),
        )
        employee_pk = int(cur.lastrowid)
        con.execute(
            """INSERT INTO employee_position_history(employee_pk, position_id, effective_start, effective_end, source_note,
                   is_current, change_type, change_reason, remarks, changed_by, changed_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP)""",
            (employee_pk, payload.position_id, iso(payload.join_date) or today().isoformat(), None, "Manual entry v22.2", 1,
             "Initial Position", "New Employee", payload.remarks, user["user_id"]),
        )
        con.commit()
    write_audit(user, "CREATE_EMPLOYEE", "Employee Master", code, f"employee_pk={employee_pk}; position_id={payload.position_id}; state=Pending Initial Assessment")
    return {"status":"success","employee_pk":employee_pk,"assessment_state":"Pending Initial Assessment"}


@app.put("/api/employees/{employee_pk}")
def update_employee(employee_pk: int, payload: EmployeeUpdate) -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        require_row(con, "employees", "employee_pk", employee_pk, "Employee")
        con.execute(
            """UPDATE employees SET employee_name=?, department=?, join_date=?, employment_status=?, exit_date=?,
                   status_reason=?, remarks=?, updated_by=?, updated_at=CURRENT_TIMESTAMP WHERE employee_pk=?""",
            (payload.employee_name.strip(), payload.department, iso(payload.join_date), payload.employment_status,
             iso(payload.exit_date), payload.status_reason, payload.remarks, user["user_id"], employee_pk),
        )
        con.commit()
    write_audit(user, "UPDATE_EMPLOYEE", "Employee Master", str(employee_pk), f"name={payload.employee_name}; status={payload.employment_status}")
    return {"status":"success","employee_pk":employee_pk}


@app.post("/api/employees/{employee_pk}/position-change")
def change_employee_position(employee_pk: int, payload: PositionChange) -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        con.execute("PRAGMA foreign_keys=ON")
        emp = require_row(con, "employees", "employee_pk", employee_pk, "Employee")
        pos = require_row(con, "positions", "position_id", payload.position_id, "Position")
        old_position = emp["current_position_id"]
        if old_position == payload.position_id:
            raise HTTPException(status_code=400, detail="Position baru sama dengan position saat ini.")
        con.execute(
            """UPDATE employee_position_history SET effective_end=date(?, '-1 day'), is_current=0
               WHERE employee_pk=? AND (is_current=1 OR effective_end IS NULL)""",
            (iso(payload.effective_date), employee_pk),
        )
        con.execute(
            """INSERT INTO employee_position_history(employee_pk,position_id,effective_start,is_current,change_type,change_reason,remarks,changed_by,changed_at,source_note)
               VALUES (?,?,?,1,'Position Change',?,?,?,CURRENT_TIMESTAMP,'Manual change v22.2')""",
            (employee_pk, payload.position_id, iso(payload.effective_date), payload.change_reason, payload.remarks, user["user_id"]),
        )
        old_state = emp["assessment_state"] or "Assessed"
        if old_state in {"Pending Initial Assessment", "Assessment Ready"}:
            new_state = "Pending Initial Assessment"
            con.execute(
                """UPDATE employees SET current_position_id=?, department=COALESCE(?,department), assessment_state=?,
                       initial_data_confirmed=0, initial_data_confirmed_at=NULL, initial_data_confirmed_by=NULL,
                       updated_by=?, updated_at=CURRENT_TIMESTAMP WHERE employee_pk=?""",
                (payload.position_id, pos["department"], new_state, user["user_id"], employee_pk),
            )
        else:
            new_state = "Pending Reassessment"
            con.execute(
                """UPDATE employees SET current_position_id=?, department=COALESCE(?,department), assessment_state=?,
                       updated_by=?, updated_at=CURRENT_TIMESTAMP WHERE employee_pk=?""",
                (payload.position_id, pos["department"], new_state, user["user_id"], employee_pk),
            )
        con.commit()
    write_audit(user, "CHANGE_POSITION", "Employee Master", str(employee_pk), f"from={old_position}; to={payload.position_id}; effective={payload.effective_date}; state={new_state}")
    return {"status":"success","employee_pk":employee_pk,"assessment_state":new_state}


@app.post("/api/employees/{employee_pk}/deactivate")
def deactivate_employee(employee_pk: int, payload: EmployeeDeactivate) -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        require_row(con, "employees", "employee_pk", employee_pk, "Employee")
        con.execute(
            """UPDATE employees SET active_flag=0, employment_status=?, exit_date=?, status_reason=?, remarks=COALESCE(?,remarks),
                   assessment_state='Archived', record_status='Archived', updated_by=?, updated_at=CURRENT_TIMESTAMP WHERE employee_pk=?""",
            (payload.employment_status, iso(payload.effective_date), payload.reason, payload.remarks, user["user_id"], employee_pk),
        )
        con.execute(
            """UPDATE employee_position_history SET effective_end=COALESCE(effective_end,?), is_current=0
               WHERE employee_pk=? AND (is_current=1 OR effective_end IS NULL)""",
            (iso(payload.effective_date), employee_pk),
        )
        con.commit()
    write_audit(user, "DEACTIVATE_EMPLOYEE", "Employee Master", str(employee_pk), f"status={payload.employment_status}; reason={payload.reason}")
    return {"status":"success","employee_pk":employee_pk,"employment_status":payload.employment_status}


@app.get("/api/employees/{employee_pk}/purge-impact")
def employee_purge_impact(employee_pk: int) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        employee = require_row(con, "employees", "employee_pk", employee_pk, "Employee")
        impact = calculate_employee_purge_impact(con, employee_pk)
    return {
        "employee_pk": employee_pk,
        "employee_code": employee["employee_code"],
        "employee_name": employee["employee_name"],
        "eligible": is_test_employee_code(str(employee["employee_code"] or "")),
        "impact": impact,
        "related_records": sum(impact.values()),
    }


@app.delete("/api/employees/{employee_pk}/purge")
def purge_test_employee(employee_pk: int, payload: EmployeePurge) -> dict[str, Any]:
    user = current_user()
    if user.get("role") != "supervisor_tcd":
        raise HTTPException(status_code=403, detail="Hanya Supervisor TCD yang dapat menghapus employee test.")
    reason = payload.reason.strip()
    if not reason:
        raise HTTPException(status_code=400, detail="Deletion Reason wajib diisi.")

    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        employee = require_row(con, "employees", "employee_pk", employee_pk, "Employee")
        employee_code = str(employee["employee_code"] or "").strip()
        employee_name = str(employee["employee_name"] or "").strip()
        if not is_test_employee_code(employee_code):
            raise HTTPException(
                status_code=409,
                detail="Permanent delete hanya tersedia untuk Employee ID berawalan UAT-, TEST-, atau DUMMY-. Gunakan Deactivate untuk employee perusahaan.",
            )
        if payload.confirmation_code.strip().casefold() != employee_code.casefold():
            raise HTTPException(status_code=400, detail=f"Ketik Employee ID {employee_code} untuk mengonfirmasi penghapusan.")
        impact = calculate_employee_purge_impact(con, employee_pk)
        evidence_rows = con.execute(
            "SELECT document_id,file_path FROM evidence_documents WHERE employee_pk=?",
            (employee_pk,),
        ).fetchall()

    # A complete database + evidence backup is mandatory before destructive UAT cleanup.
    safety = create_backup(f"pre_purge_employee_{employee_pk}", user["user_id"])
    target = f"{employee_code} — {employee_name}"

    with sqlite3.connect(DB_PATH) as con:
        con.execute("PRAGMA foreign_keys=ON")
        try:
            deleted = purge_employee_records(con, employee_pk)
            refresh_notifications(con)
            fk_issues = con.execute("PRAGMA foreign_key_check").fetchall()
            if fk_issues:
                raise RuntimeError(f"Foreign key check gagal: {len(fk_issues)} issue.")
            con.execute(
                "INSERT INTO audit_log(user_id,action,dataset_type,target,details) VALUES (?,?,?,?,?)",
                (
                    user["user_id"], "PURGE_TEST_EMPLOYEE", "Employee Master", target,
                    f"employee_pk={employee_pk}; reason={reason}; backup={safety['filename']}; records={sum(deleted.values())}; evidence_candidates={len(evidence_rows)}",
                ),
            )
            con.commit()
        except Exception:
            con.rollback()
            raise

    # Active evidence storage is cleaned without destroying the recovery copy.
    quarantine = ROOT / "backups" / "purged_evidence" / f"employee_{employee_pk}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    moved_files = 0
    move_warnings: list[str] = []
    evidence_root = (ROOT / "evidence").resolve()
    for row in evidence_rows:
        source = (ROOT / str(row["file_path"] or "")).resolve()
        if evidence_root not in source.parents or not source.exists():
            continue
        try:
            quarantine.mkdir(parents=True, exist_ok=True)
            shutil.move(str(source), str(quarantine / source.name))
            moved_files += 1
        except Exception as exc:
            move_warnings.append(f"document_id={row['document_id']}: {exc}")

    # Drop any in-memory conversational reference to the deleted entity.
    for key, value in list(_AI_CONTEXTS.items()):
        if employee_pk in set(value.get("employee_pks") or []):
            _AI_CONTEXTS.pop(key, None)

    return {
        "status": "success",
        "employee_pk": employee_pk,
        "employee_code": employee_code,
        "employee_name": employee_name,
        "deleted_records": deleted,
        "impact_before_delete": impact,
        "safety_backup": safety["filename"],
        "evidence_quarantined": moved_files,
        "warnings": move_warnings,
    }


@app.get("/api/employees/{employee_pk}/assessment-readiness")
def employee_assessment_readiness(employee_pk: int) -> dict[str, Any]:
    try:
        with sqlite3.connect(DB_PATH) as con:
            return assessment_readiness(con, employee_pk)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc))


@app.post("/api/employees/{employee_pk}/assessment/confirm-initial-data")
def confirm_initial_assessment_data(employee_pk: int, payload: AssessmentConfirm) -> dict[str, Any]:
    user = current_user()
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.row_factory = sqlite3.Row
            ready = assessment_readiness(con, employee_pk)
            if not ready["active"]:
                raise HTTPException(status_code=409, detail="Employee tidak aktif.")
            if ready["assessment_state"] not in {"Pending Initial Assessment", "Assessment Ready"}:
                raise HTTPException(status_code=409, detail=f"Initial data confirmation tidak berlaku untuk state {ready['assessment_state']}.")
            if not ready["can_confirm_initial"]:
                raise HTTPException(status_code=409, detail="Current position atau Training Requirement aktif belum tersedia.")
            con.execute(
                """UPDATE employees SET initial_data_confirmed=1,initial_data_confirmed_at=CURRENT_TIMESTAMP,
                       initial_data_confirmed_by=?,assessment_state='Assessment Ready',updated_by=?,updated_at=CURRENT_TIMESTAMP
                   WHERE employee_pk=?""",
                (user["user_id"], user["user_id"], employee_pk),
            )
            con.commit()
        write_audit(user, "CONFIRM_INITIAL_DATA", "Assessment", str(employee_pk),
                    f"requirements={ready['applicable_requirements']}; training={ready['training_history_records']}; certifications={ready['certification_records']}; remarks={payload.remarks or ''}")
        return {"status":"success","employee_pk":employee_pk,"assessment_state":"Assessment Ready","readiness":ready}
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc))


@app.post("/api/employees/{employee_pk}/assessment/run")
def run_initial_assessment(employee_pk: int) -> dict[str, Any]:
    user = current_user()
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.execute("PRAGMA foreign_keys=ON")
            result = run_employee_assessment(con, employee_pk, "Initial Assessment", user["user_id"])
            con.commit()
        write_audit(user, "RUN_INITIAL_ASSESSMENT", "Assessment", str(employee_pk),
                    f"run_id={result['run_id']}; requirements={result['applicable_requirements']}; distribution={result['distribution']}")
        return {"status":"success", **result}
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc))


@app.post("/api/employees/{employee_pk}/assessment/reassess")
def run_reassessment(employee_pk: int) -> dict[str, Any]:
    user = current_user()
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.execute("PRAGMA foreign_keys=ON")
            result = run_employee_assessment(con, employee_pk, "Reassessment", user["user_id"])
            con.commit()
        write_audit(user, "RUN_REASSESSMENT", "Assessment", str(employee_pk),
                    f"run_id={result['run_id']}; requirements={result['applicable_requirements']}; distribution={result['distribution']}")
        return {"status":"success", **result}
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc))


@app.get("/api/assessment-queue")
def assessment_queue(include_assessed: bool = False, limit: int = 5000) -> dict[str, Any]:
    where = "" if include_assessed else "AND COALESCE(e.assessment_state,'Assessed') <> 'Assessed'"
    df = query_df(f"""
        SELECT e.employee_pk,e.employee_code,e.employee_name,p.position_name,
               COALESCE(e.assessment_state,'Assessed') assessment_state,e.initial_data_confirmed,
               e.initial_data_confirmed_at,e.last_assessment_at,e.active_flag,
               (SELECT COUNT(*) FROM training_requirements tr
                 WHERE tr.active_flag=1 AND COALESCE(tr.requirement_status,'Active')='Active'
                   AND COALESCE(tr.position_standard_id,tr.position_id)=e.current_position_id
                   AND (tr.effective_from IS NULL OR date(tr.effective_from)<=date('now'))
                   AND (tr.effective_to IS NULL OR date(tr.effective_to)>=date('now'))) applicable_requirements,
               (SELECT COUNT(*) FROM training_history th WHERE th.employee_pk=e.employee_pk AND COALESCE(th.record_status,'Active')<>'Archived') training_history_records,
               (SELECT COUNT(*) FROM certifications c WHERE c.employee_pk=e.employee_pk AND COALESCE(c.record_status,'Active')<>'Archived') certification_records
        FROM employees e LEFT JOIN positions p ON p.position_id=e.current_position_id
        WHERE e.active_flag=1 {where}
        ORDER BY CASE COALESCE(e.assessment_state,'Assessed')
                   WHEN 'Pending Initial Assessment' THEN 1 WHEN 'Assessment Ready' THEN 2
                   WHEN 'Pending Reassessment' THEN 3 ELSE 9 END, e.employee_name
        LIMIT ?
    """, (min(max(limit,1),5000),)).fillna("")
    summary = query_df("""
        SELECT COALESCE(assessment_state,'Assessed') state, COUNT(*) total
        FROM employees WHERE active_flag=1 GROUP BY COALESCE(assessment_state,'Assessed')
    """).fillna("").to_dict(orient="records")
    return {"summary":summary,"records":df.to_dict(orient="records")}


@app.get("/api/certifications")
def certifications(status: str | None = None, search: str | None = None, limit: int = 2000, include_history: bool = False) -> list[dict[str, Any]]:
    limit = max(1, min(limit, 5000))
    where = "" if include_history else "WHERE COALESCE(c.record_status,'Active') <> 'Archived'"
    rows = query_df(f"""
        SELECT c.certification_id, c.employee_pk, COALESCE(e.employee_name,c.employee_name_raw) employee_name,
               p.position_name, COALESCE(cc.certification_name,c.certification_name_raw) certification_name,
               c.certification_catalog_id, c.certificate_number, COALESCE(c.issuer,c.issuer_type) issuer,
               COALESCE(c.issue_date,c.certification_date) issue_date, COALESCE(c.expiry_date,c.expired_date) expiry_date,
               COALESCE(c.certification_status,c.status) source_status, c.renewal_status, c.previous_record_id,
               c.reason_code, COALESCE(c.remarks,c.note) remarks, c.record_status
        FROM certifications c
        LEFT JOIN employees e ON e.employee_pk=c.employee_pk
        LEFT JOIN positions p ON p.position_id=e.current_position_id
        LEFT JOIN certification_catalog cc ON cc.certification_catalog_id=c.certification_catalog_id
        {where}
    """)
    if rows.empty:
        return []
    rows["expiry_dt"] = pd.to_datetime(rows["expiry_date"], errors="coerce")
    d = pd.Timestamp(today())
    def display(r):
        src=str(r.get("source_status","") or "")
        if src in {"Superseded","Inactive","Not Renewed"}: return src
        x=r["expiry_dt"]
        if pd.notna(x) and x < d: return "Expired"
        if pd.notna(x) and d <= x <= d+pd.Timedelta(days=90): return "Near Expiry"
        if pd.notna(x) and x > d: return "Active"
        if src.lower() in {"dalam proses","in process"}: return "In Process"
        return src or "Not Available"
    rows["display_status"] = rows.apply(display, axis=1)
    rows["renewal_status"] = rows.apply(lambda r: renewal_status_for(r.get("display_status"), r.get("renewal_status")), axis=1)
    if status:
        rows=rows[rows.display_status.str.lower()==status.lower()]
    if search:
        q=search.lower(); mask=rows[["employee_name","position_name","certification_name","certificate_number"]].fillna("").astype(str).apply(lambda c:c.str.lower().str.contains(q,regex=False)).any(axis=1); rows=rows[mask]
    rows=rows.sort_values(["expiry_dt","employee_name"],na_position="last")
    cols=["certification_id","employee_pk","employee_name","position_name","certification_name","certification_catalog_id","certificate_number","issuer","issue_date","expiry_date","display_status","source_status","renewal_status","previous_record_id","reason_code","remarks","record_status"]
    return rows[cols].fillna("").head(limit).to_dict(orient="records")


@app.post("/api/certifications")
def create_certification(payload: CertificationCreate) -> dict[str, Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        con.execute("PRAGMA foreign_keys=ON")
        require_row(con,"employees","employee_pk",payload.employee_pk,"Employee")
        catalog_id=resolve_certification_catalog(con,payload.certification_name,payload.issuer)
        status=certification_status_for(payload.expiry_date,payload.certification_status)
        cur=con.execute("""INSERT INTO certifications(employee_pk,certification_catalog_id,certification_name_raw,certificate_number,issuer,issuer_type,issue_date,certification_date,expiry_date,expired_date,status,certification_status,renewal_status,remarks,note,record_status,created_by,updated_by,created_at,updated_at)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)""",
                        (payload.employee_pk,catalog_id,payload.certification_name.strip(),payload.certificate_number,payload.issuer,payload.issuer,iso(payload.issue_date),iso(payload.issue_date),iso(payload.expiry_date),iso(payload.expiry_date),status,status,renewal_status_for(status),payload.remarks,payload.remarks,"Active",user["user_id"],user["user_id"]))
        cid=int(cur.lastrowid); mark_employee_for_reassessment(con,payload.employee_pk,user["user_id"]); state=con.execute("SELECT assessment_state FROM employees WHERE employee_pk=?",(payload.employee_pk,)).fetchone()[0]; con.commit()
    write_audit(user,"CREATE_CERTIFICATION","Certification",str(cid),f"employee_pk={payload.employee_pk}; certification={payload.certification_name}; status={status}")
    return {"status":"success","certification_id":cid,"assessment_state":state}


@app.put("/api/certifications/{certification_id}")
def update_certification(certification_id:int,payload:CertificationUpdate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"certifications","certification_id",certification_id,"Certification")
        require_row(con,"employees","employee_pk",payload.employee_pk,"Employee")
        catalog_id=resolve_certification_catalog(con,payload.certification_name,payload.issuer)
        status=certification_status_for(payload.expiry_date,payload.certification_status)
        con.execute("""UPDATE certifications SET employee_pk=?,certification_catalog_id=?,certification_name_raw=?,certificate_number=?,issuer=?,issuer_type=?,issue_date=?,certification_date=?,expiry_date=?,expired_date=?,status=?,certification_status=?,renewal_status=?,reason_code=?,remarks=?,note=?,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE certification_id=?""",
                    (payload.employee_pk,catalog_id,payload.certification_name.strip(),payload.certificate_number,payload.issuer,payload.issuer,iso(payload.issue_date),iso(payload.issue_date),iso(payload.expiry_date),iso(payload.expiry_date),status,status,renewal_status_for(status,payload.renewal_status),payload.reason_code,payload.remarks,payload.remarks,user["user_id"],certification_id))
        mark_employee_for_reassessment(con,int(old["employee_pk"]),user["user_id"]); mark_employee_for_reassessment(con,payload.employee_pk,user["user_id"]); con.commit()
    write_audit(user,"UPDATE_CERTIFICATION","Certification",str(certification_id),f"status={status}; certification={payload.certification_name}")
    return {"status":"success","certification_id":certification_id}


@app.post("/api/certifications/{certification_id}/renew")
def renew_certification(certification_id:int,payload:CertificationRenew)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"certifications","certification_id",certification_id,"Certification")
        if str(old["record_status"] or "Active") == "Archived":
            raise HTTPException(status_code=400,detail="Record sertifikasi lama sudah diarsipkan.")
        catalog_id=old["certification_catalog_id"]
        cert_name=old["certification_name_raw"]
        if catalog_id:
            row=con.execute("SELECT certification_name FROM certification_catalog WHERE certification_catalog_id=?",(catalog_id,)).fetchone(); cert_name=(row[0] if row else cert_name)
        new_status=certification_status_for(payload.expiry_date,"Auto")
        cur=con.execute("""INSERT INTO certifications(employee_pk,certification_catalog_id,employee_code_raw,employee_name_raw,certification_name_raw,held_by,issuer_type,certificate_number,issuer,issue_date,certification_date,expiry_date,expired_date,status,certification_status,renewal_status,previous_record_id,renewed_date,remarks,note,record_status,created_by,updated_by,created_at,updated_at)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)""",
                        (old["employee_pk"],catalog_id,old["employee_code_raw"],old["employee_name_raw"],cert_name,old["held_by"],payload.issuer or old["issuer_type"],payload.certificate_number,payload.issuer or old["issuer"],iso(payload.issue_date),iso(payload.issue_date),iso(payload.expiry_date),iso(payload.expiry_date),new_status,new_status,renewal_status_for(new_status),certification_id,iso(payload.issue_date),payload.remarks,payload.remarks,"Active",user["user_id"],user["user_id"]))
        new_id=int(cur.lastrowid)
        con.execute("""UPDATE certifications SET certification_status='Superseded',status='Superseded',renewal_status='Completed',renewed_date=?,record_status='Archived',updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE certification_id=?""",(iso(payload.issue_date),user["user_id"],certification_id))
        mark_employee_for_reassessment(con,int(old["employee_pk"]),user["user_id"]); con.commit()
    write_audit(user,"RENEW_CERTIFICATION","Certification",str(certification_id),f"new_certification_id={new_id}; employee_pk={old['employee_pk']}")
    return {"status":"success","old_certification_id":certification_id,"new_certification_id":new_id,"assessment_state":"Pending Reassessment"}


@app.post("/api/certifications/{certification_id}/not-renewed")
def certification_not_renewed(certification_id:int,payload:CertificationNotRenewed)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"certifications","certification_id",certification_id,"Certification")
        con.execute("UPDATE certifications SET renewal_status='Not Renewed',reason_code=?,remarks=COALESCE(?,remarks),updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE certification_id=?",(payload.reason_code,payload.remarks,user["user_id"],certification_id)); con.commit()
    write_audit(user,"MARK_CERTIFICATION_NOT_RENEWED","Certification",str(certification_id),f"reason={payload.reason_code}")
    return {"status":"success","certification_id":certification_id}


@app.post("/api/training-catalog")
def create_training_catalog(payload:TrainingCatalogCreate)->dict[str,Any]:
    user=current_user(); name=payload.training_name.strip()
    if not name: raise HTTPException(status_code=400,detail="Nama training wajib diisi.")
    with sqlite3.connect(DB_PATH) as con:
        row=con.execute("SELECT training_id FROM training_catalog WHERE lower(training_name)=lower(?)",(name,)).fetchone()
        if row: raise HTTPException(status_code=409,detail="Training sudah ada di catalog.")
        cur=con.execute("INSERT INTO training_catalog(training_name,category,active_flag) VALUES (?,?,1)",(name,payload.category)); tid=int(cur.lastrowid); con.commit()
    write_audit(user,"CREATE_TRAINING_CATALOG","Training Catalog",str(tid),name)
    return {"status":"success","training_id":tid}


@app.put("/api/training-catalog/{training_id}")
def update_training_catalog(training_id:int,payload:TrainingCatalogUpdate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        require_row(con,"training_catalog","training_id",training_id,"Training"); con.execute("UPDATE training_catalog SET training_name=?,category=?,active_flag=? WHERE training_id=?",(payload.training_name.strip(),payload.category,payload.active_flag,training_id)); con.commit()
    write_audit(user,"UPDATE_TRAINING_CATALOG","Training Catalog",str(training_id),payload.training_name)
    return {"status":"success","training_id":training_id}


@app.get("/api/training-requirements")
def training_requirements(position: str | None = None, employee: str | None = None, employee_pk: int | None = None, department: str | None = None, requirement_type: str | None = None, search: str | None = None, limit: int = 2000) -> list[dict[str, Any]]:
    limit = max(1, min(limit, 5000))
    sql = """
        SELECT tr.requirement_id, tr.training_id, COALESCE(tr.position_standard_id,tr.position_id) AS position_id,
               COALESCE(pstd.position_name, p.position_name, tr.position_raw) AS position_name,
               COALESCE(pstd.department, p.department, GROUP_CONCAT(DISTINCT pemp.department), '') AS department,
               tc.training_name,
               tr.requirement_type,
               tr.delivery_type,
               tr.regulatory_flag,
               tr.requirement_code,
               tr.frequency_type,
               tr.frequency_value,
               tr.frequency_unit, tr.effective_from, tr.effective_to,
               COALESCE(tr.requirement_status, CASE WHEN tr.active_flag=1 THEN 'Active' ELSE 'Inactive' END) AS requirement_status,
               tr.requirement_source, tr.remarks, tr.active_flag,
               COUNT(DISTINCT e.employee_pk) AS employee_count,
               GROUP_CONCAT(DISTINCT e.employee_name) AS employee_names
        FROM training_requirements tr
        JOIN training_catalog tc ON tc.training_id = tr.training_id
        LEFT JOIN positions p ON p.position_id = tr.position_id
        LEFT JOIN positions pstd ON pstd.position_id = tr.position_standard_id
        LEFT JOIN position_crosswalk pc ON lower(pc.source_position_key)=lower(coalesce(tr.position_key,'')) AND pc.crosswalk_status='Exact Match'
        LEFT JOIN employees e ON e.active_flag = 1 AND e.current_position_id = COALESCE(tr.position_standard_id, pc.standard_position_id, tr.position_id)
        LEFT JOIN positions pemp ON pemp.position_id = e.current_position_id
        GROUP BY tr.requirement_id
        ORDER BY position_name, tc.training_name
    """
    rows = query_df(sql)
    if rows.empty:
        return []
    for col in ["position_name","department","training_name","requirement_type","delivery_type","regulatory_flag","employee_names"]:
        rows[col] = rows[col].fillna("").astype(str)
    # Employee filter is an employee context, not a position aggregate.
    # Resolve the selected employee first, then return only the requirements
    # applicable to that employee's current position and label rows with that
    # employee only. This keeps screen and Export Filtered semantically aligned.
    selected_employee = None
    if employee_pk is not None or employee:
        if employee_pk is not None:
            emp = query_df("""
                SELECT e.employee_pk,e.employee_code,e.employee_name,e.current_position_id,
                       p.position_name,COALESCE(e.department,p.department,'') AS department
                FROM employees e LEFT JOIN positions p ON p.position_id=e.current_position_id
                WHERE e.active_flag=1 AND e.employee_pk=?
            """, (employee_pk,))
        else:
            emp = query_df("""
                SELECT e.employee_pk,e.employee_code,e.employee_name,e.current_position_id,
                       p.position_name,COALESCE(e.department,p.department,'') AS department
                FROM employees e LEFT JOIN positions p ON p.position_id=e.current_position_id
                WHERE e.active_flag=1 AND lower(e.employee_name)=lower(?)
            """, (employee,))
        if not emp.empty:
            selected_employee = emp.iloc[0]
            pos_id = int(selected_employee["current_position_id"] or 0)
            if pos_id:
                rows = rows[pd.to_numeric(rows["position_id"], errors="coerce").fillna(-1).astype(int) == pos_id].copy()
            else:
                rows = rows.iloc[0:0].copy()
            if not rows.empty:
                rows["employee_pk"] = int(selected_employee["employee_pk"])
                rows["employee_code"] = str(selected_employee["employee_code"] or "")
                rows["employee_count"] = 1
                rows["employee_names"] = str(selected_employee["employee_name"] or "")
                rows["department"] = str(selected_employee["department"] or rows["department"].iloc[0] or "")

    if position:
        q = position.lower()
        rows = rows[rows.position_name.str.lower().str.contains(q, regex=False)]
    if department:
        q = department.lower()
        rows = rows[rows.department.str.lower().str.contains(q, regex=False)]
    if requirement_type:
        q = requirement_type.lower()
        rows = rows[rows.requirement_type.str.lower() == q]
    if search:
        q = search.lower()
        mask = rows[["position_name","department","training_name","employee_names","requirement_type"]].apply(
            lambda c: c.str.lower().str.contains(q, regex=False)
        ).any(axis=1)
        rows = rows[mask]
    if (employee_pk is not None or employee) and rows.empty and selected_employee is not None:
        rows = pd.DataFrame([{
            "requirement_id":"",
            "position_id": selected_employee["current_position_id"] or "",
            "position_name":selected_employee["position_name"] or "—",
            "department":selected_employee["department"] or "—",
            "training_name":"Belum ada Training Requirement terpetakan",
            "requirement_type":"Data Not Available",
            "delivery_type":"—",
            "regulatory_flag":"—",
            "requirement_code":"—",
            "frequency_type":"", "frequency_value":"", "frequency_unit":"",
            "employee_pk":int(selected_employee["employee_pk"]),
            "employee_code":str(selected_employee["employee_code"] or ""),
            "employee_count":1, "employee_names":selected_employee["employee_name"]
        }])
    rows["employee_names"] = rows.apply(
        lambda r: r["employee_names"] if r["employee_names"] else "Belum ada karyawan yang terpetakan pada posisi ini", axis=1
    )
    return rows.head(limit).fillna("").to_dict(orient="records")


@app.post("/api/training-requirements")
def create_training_requirement(payload:TrainingRequirementCreate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        pos=require_row(con,"positions","position_id",payload.position_id,"Position"); trn=require_row(con,"training_catalog","training_id",payload.training_id,"Training")
        dup=con.execute("SELECT requirement_id FROM training_requirements WHERE COALESCE(position_standard_id,position_id)=? AND training_id=? AND COALESCE(requirement_status,'Active')='Active'",(payload.position_id,payload.training_id)).fetchone()
        if dup: raise HTTPException(status_code=409,detail=f"Requirement aktif yang sama sudah ada (ID {dup[0]}).")
        code=payload.requirement_code or ("★" if payload.requirement_type.lower()=="mandatory" else "")
        cur=con.execute("""INSERT INTO training_requirements(training_id,position_id,position_standard_id,position_raw,position_key,requirement_code,requirement_type,delivery_type,regulatory_flag,active_flag,effective_from,effective_to,requirement_status,requirement_source,remarks,updated_by,updated_at) VALUES (?,?,?,?,?,?,?,?,?,1,?,?,?,?,?,?,CURRENT_TIMESTAMP)""",
                        (payload.training_id,payload.position_id,payload.position_id,pos["position_name"],str(pos["position_name"] or "").lower(),code,payload.requirement_type,payload.delivery_type,payload.regulatory_flag,iso(payload.effective_from),iso(payload.effective_to),"Active",payload.requirement_source,payload.remarks,user["user_id"]))
        rid=int(cur.lastrowid); affected=mark_position_for_reassessment(con,payload.position_id,user["user_id"]); con.commit()
    write_audit(user,"CREATE_TRAINING_REQUIREMENT","Training Requirement",str(rid),f"position={pos['position_name']}; training={trn['training_name']}; affected={affected}")
    return {"status":"success","requirement_id":rid,"employees_marked_for_reassessment":affected}


@app.put("/api/training-requirements/{requirement_id}")
def update_training_requirement(requirement_id:int,payload:TrainingRequirementUpdate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"training_requirements","requirement_id",requirement_id,"Training Requirement"); pos=require_row(con,"positions","position_id",payload.position_id,"Position"); require_row(con,"training_catalog","training_id",payload.training_id,"Training")
        code=payload.requirement_code or ("★" if payload.requirement_type.lower()=="mandatory" else "")
        con.execute("""UPDATE training_requirements SET training_id=?,position_id=?,position_standard_id=?,position_raw=?,position_key=?,requirement_code=?,requirement_type=?,delivery_type=?,regulatory_flag=?,active_flag=?,effective_from=?,effective_to=?,requirement_status=?,requirement_source=?,remarks=?,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE requirement_id=?""",
                    (payload.training_id,payload.position_id,payload.position_id,pos["position_name"],str(pos["position_name"] or "").lower(),code,payload.requirement_type,payload.delivery_type,payload.regulatory_flag,1 if payload.requirement_status=='Active' else 0,iso(payload.effective_from),iso(payload.effective_to),payload.requirement_status,payload.requirement_source,payload.remarks,user["user_id"],requirement_id))
        affected=mark_position_for_reassessment(con,int(old["position_standard_id"] or old["position_id"] or payload.position_id),user["user_id"]); affected+=mark_position_for_reassessment(con,payload.position_id,user["user_id"]); con.commit()
    write_audit(user,"UPDATE_TRAINING_REQUIREMENT","Training Requirement",str(requirement_id),f"status={payload.requirement_status}; affected={affected}")
    return {"status":"success","requirement_id":requirement_id,"employees_marked_for_reassessment":affected}


@app.post("/api/training-requirements/{requirement_id}/deactivate")
def deactivate_training_requirement(requirement_id:int,payload:RequirementDeactivate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"training_requirements","requirement_id",requirement_id,"Training Requirement"); position_id=int(old["position_standard_id"] or old["position_id"] or 0)
        con.execute("UPDATE training_requirements SET active_flag=0,requirement_status='Inactive',effective_to=?,reason_code=?,remarks=COALESCE(?,remarks),updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE requirement_id=?",(iso(payload.effective_to),payload.reason_code,payload.remarks,user["user_id"],requirement_id)); affected=mark_position_for_reassessment(con,position_id,user["user_id"]) if position_id else 0; con.commit()
    write_audit(user,"DEACTIVATE_TRAINING_REQUIREMENT","Training Requirement",str(requirement_id),f"reason={payload.reason_code}; affected={affected}")
    return {"status":"success","requirement_id":requirement_id,"employees_marked_for_reassessment":affected}


@app.get("/api/training-history")
def training_history(employee: str | None = None, search: str | None = None, limit: int = 2000, include_archived: bool = False) -> list[dict[str, Any]]:
    limit=max(1,min(limit,5000)); where="" if include_archived else "WHERE COALESCE(th.record_status,'Active') <> 'Archived'"
    df=query_df(f"""SELECT th.history_id,th.employee_pk,e.employee_name,p.position_name,COALESCE(th.training_id,th.training_pk) training_id,COALESCE(th.training_name,th.training_name_raw,tc.training_name) training_name,th.training_date,th.completion_date,th.completion_year,COALESCE(th.result_status,th.completion_status) completion_status,th.certificate_reference,th.provider,th.valid_until,th.evidence_status,COALESCE(th.evidence_note,th.note) evidence_note,th.remarks,th.record_status FROM training_history th JOIN employees e ON e.employee_pk=th.employee_pk LEFT JOIN positions p ON p.position_id=e.current_position_id LEFT JOIN training_catalog tc ON tc.training_id=COALESCE(th.training_pk,th.training_id) {where} ORDER BY th.completion_date DESC,e.employee_name LIMIT ?""",(limit,)).fillna("")
    if employee or search:
        q=(employee or search or "").lower(); df=df[df[["employee_name","position_name","training_name"]].astype(str).apply(lambda c:c.str.lower().str.contains(q,regex=False)).any(axis=1)]
    return df.to_dict(orient="records")


@app.post("/api/training-history")
def create_training_history(payload:TrainingHistoryCreate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        require_row(con,"employees","employee_pk",payload.employee_pk,"Employee"); trn=require_row(con,"training_catalog","training_id",payload.training_id,"Training")
        completion=payload.completion_date or payload.training_date
        cur=con.execute("""INSERT INTO training_history(employee_pk,training_id,training_pk,training_name_raw,training_name,training_date,completion_date,completion_year,completion_status,result_status,certificate_reference,provider,evidence_note,note,valid_until,evidence_status,record_status,remarks,created_by,updated_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)""",
                        (payload.employee_pk,payload.training_id,payload.training_id,trn["training_name"],trn["training_name"],iso(payload.training_date),iso(completion),completion.year if completion else None,payload.result_status,payload.result_status,payload.certificate_reference,payload.provider,payload.evidence_note,payload.evidence_note,iso(payload.valid_until),"Available" if (payload.evidence_note or payload.certificate_reference) else "No attachment","Active",payload.remarks,user["user_id"],user["user_id"]))
        hid=int(cur.lastrowid); mark_employee_for_reassessment(con,payload.employee_pk,user["user_id"]); state=con.execute("SELECT assessment_state FROM employees WHERE employee_pk=?",(payload.employee_pk,)).fetchone()[0]; con.commit()
    write_audit(user,"CREATE_TRAINING_HISTORY","Training History",str(hid),f"employee_pk={payload.employee_pk}; training={trn['training_name']}; result={payload.result_status}")
    return {"status":"success","history_id":hid,"assessment_state":state}


@app.put("/api/training-history/{history_id}")
def update_training_history(history_id:int,payload:TrainingHistoryUpdate)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"training_history","history_id",history_id,"Training History"); require_row(con,"employees","employee_pk",payload.employee_pk,"Employee"); trn=require_row(con,"training_catalog","training_id",payload.training_id,"Training"); completion=payload.completion_date or payload.training_date
        con.execute("""UPDATE training_history SET employee_pk=?,training_id=?,training_pk=?,training_name_raw=?,training_name=?,training_date=?,completion_date=?,completion_year=?,completion_status=?,result_status=?,certificate_reference=?,provider=?,evidence_note=?,note=?,valid_until=?,evidence_status=?,remarks=?,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE history_id=?""",
                    (payload.employee_pk,payload.training_id,payload.training_id,trn["training_name"],trn["training_name"],iso(payload.training_date),iso(completion),completion.year if completion else None,payload.result_status,payload.result_status,payload.certificate_reference,payload.provider,payload.evidence_note,payload.evidence_note,iso(payload.valid_until),"Available" if (payload.evidence_note or payload.certificate_reference) else "No attachment",payload.remarks,user["user_id"],history_id))
        mark_employee_for_reassessment(con,int(old["employee_pk"]),user["user_id"]); mark_employee_for_reassessment(con,payload.employee_pk,user["user_id"]); con.commit()
    write_audit(user,"UPDATE_TRAINING_HISTORY","Training History",str(history_id),f"result={payload.result_status}")
    return {"status":"success","history_id":history_id}


@app.post("/api/training-history/{history_id}/archive")
def archive_training_history(history_id:int)->dict[str,Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        old=require_row(con,"training_history","history_id",history_id,"Training History"); con.execute("UPDATE training_history SET record_status='Archived',updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE history_id=?",(user["user_id"],history_id)); mark_employee_for_reassessment(con,int(old["employee_pk"]),user["user_id"]); con.commit()
    write_audit(user,"ARCHIVE_TRAINING_HISTORY","Training History",str(history_id),"soft delete")
    return {"status":"success","history_id":history_id}


@app.get("/api/competency")
def competency(status: str | None = None, search: str | None = None, limit: int = 2000) -> list[dict[str, Any]]:
    limit = max(1, min(limit, 5000))
    # Coverage view is the canonical source because it contains the 4-way coverage status,
    # including Potential Gap and Covered, and aligns with the AI Assistant.
    sql = """
        SELECT ca.employee_pk,ca.requirement_id,e.employee_name,ca.position_tna AS position_name,ca.training_name,
               ca.coverage_status AS final_status,ca.priority,ca.recommendation AS decision_action,
               COALESCE(ca.rationale,ca.recommendation) AS rationale,
               COALESCE(ca.assessment_source,da.assessment_source,'Legacy') AS assessment_source,
               COALESCE(e.assessment_state,'Assessed') assessment_state,
               COALESCE(NULLIF(da.evidence_confidence,''),
                   CASE
                       WHEN ca.coverage_status = 'Covered' THEN 'High'
                       WHEN ca.coverage_status IN ('Potential Gap','Validation Required','Expired') THEN 'Medium'
                       ELSE 'Low'
                   END) AS evidence_confidence,
               COALESCE(da.engine_version,'v21/v22.3 legacy') AS engine_version
        FROM competency_assessment ca
        JOIN employees e ON e.employee_pk=ca.employee_pk
        LEFT JOIN decision_assessment_current da
          ON da.employee_pk=ca.employee_pk AND da.requirement_id=ca.requirement_id
    """
    df = query_df(sql).fillna("")
    if status:
        df = df[df.final_status.str.lower() == status.lower()]
    if search:
        q = search.lower()
        df = df[df[["employee_name","position_name","training_name"]].astype(str).apply(lambda c: c.str.lower().str.contains(q, regex=False)).any(axis=1)]
    priority_order = {"High":1,"Validation":2,"Medium":3,"Low":4,"Not Assessed":5}
    df["_order"] = df["priority"].map(priority_order).fillna(9)
    df = df.sort_values(["_order","employee_name"])
    return df.drop(columns=["_order"]).head(limit).to_dict(orient="records")

@app.get("/api/decision-engine/status")
def decision_engine_status() -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        active_validations = con.execute("SELECT COUNT(*) FROM assessment_validations WHERE active_flag=1").fetchone()[0]
        active_overrides = con.execute("SELECT COUNT(*) FROM assessment_overrides WHERE active_flag=1").fetchone()[0]
        runs_v24 = con.execute("SELECT COUNT(*) FROM assessment_runs WHERE engine_version=?", (ENGINE_VERSION,)).fetchone()[0]
    return {
        "engine_version": ENGINE_VERSION,
        "active_validations": int(active_validations),
        "active_overrides": int(active_overrides),
        "assessment_runs_v22_4": int(runs_v24),
        "principle": "Human validates evidence/mapping; Decision Engine recalculates status. Override is recorded separately and auditable.",
    }


@app.get("/api/assessment-validations")
def list_assessment_validations(active_only: bool = True, limit: int = 1000) -> list[dict[str, Any]]:
    where = "WHERE av.active_flag=1" if active_only else ""
    df = query_df(
        f"""
        SELECT av.*, e.employee_name, tc.training_name, tr.position_id
        FROM assessment_validations av
        LEFT JOIN employees e ON e.employee_pk=av.employee_pk
        JOIN training_requirements tr ON tr.requirement_id=av.requirement_id
        JOIN training_catalog tc ON tc.training_id=tr.training_id
        {where}
        ORDER BY av.validated_at DESC, av.validation_id DESC LIMIT ?
        """,
        (min(max(limit, 1), 5000),),
    ).fillna("")
    return df.to_dict(orient="records")


@app.post("/api/assessment-validations")
def create_assessment_validation(payload: MappingValidationCreate) -> dict[str, Any]:
    user = current_user()
    scope = payload.scope.strip().title()
    if scope not in {"Employee", "Global"}:
        raise HTTPException(status_code=400, detail="Scope harus Employee atau Global.")
    decision_map = {
        "confirm equivalent": "Confirm Equivalent",
        "confirmed equivalent": "Confirm Equivalent",
        "not equivalent": "Not Equivalent",
        "need more evidence": "Need More Evidence",
        "needs more evidence": "Need More Evidence",
    }
    decision = decision_map.get(payload.decision.strip().lower())
    if not decision:
        raise HTTPException(status_code=400, detail="Decision harus Confirm Equivalent, Not Equivalent, atau Need More Evidence.")
    if scope == "Employee" and not payload.employee_pk:
        raise HTTPException(status_code=400, detail="employee_pk wajib untuk scope Employee.")

    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        req = require_row(con, "training_requirements", "requirement_id", payload.requirement_id, "Training Requirement")
        employee_pk = payload.employee_pk if scope == "Employee" else None
        if employee_pk:
            require_row(con, "employees", "employee_pk", employee_pk, "Employee")
        mapping_id = None
        if employee_pk:
            m = con.execute(
                "SELECT mapping_id FROM training_certification_mapping WHERE employee_pk=? AND requirement_id=? ORDER BY mapping_id DESC LIMIT 1",
                (employee_pk, payload.requirement_id),
            ).fetchone()
            mapping_id = int(m[0]) if m else None
        if scope == "Employee":
            con.execute(
                "UPDATE assessment_validations SET active_flag=0,revoked_by=?,revoked_at=CURRENT_TIMESTAMP WHERE requirement_id=? AND employee_pk=? AND scope='Employee' AND active_flag=1",
                (user["user_id"], payload.requirement_id, employee_pk),
            )
        else:
            con.execute(
                "UPDATE assessment_validations SET active_flag=0,revoked_by=?,revoked_at=CURRENT_TIMESTAMP WHERE requirement_id=? AND scope='Global' AND active_flag=1",
                (user["user_id"], payload.requirement_id),
            )
        cur = con.execute(
            """INSERT INTO assessment_validations(
                employee_pk,requirement_id,mapping_id,scope,decision,notes,evidence_reference,active_flag,validated_by,validated_at
            ) VALUES (?,?,?,?,?,?,?,1,?,CURRENT_TIMESTAMP)""",
            (employee_pk, payload.requirement_id, mapping_id, scope, decision, payload.notes, payload.evidence_reference, user["user_id"]),
        )
        validation_id = int(cur.lastrowid)
        if scope == "Employee":
            mark_employee_for_reassessment(con, int(employee_pk), user["user_id"])
            affected = 1
        else:
            position_id = int(req["position_standard_id"] or req["position_id"])
            affected = mark_position_for_reassessment(con, position_id, user["user_id"])
        con.commit()
    write_audit(user, "VALIDATE_MAPPING", "Assessment Validation", str(validation_id), f"scope={scope}; decision={decision}; requirement_id={payload.requirement_id}; affected={affected}")
    return {"status": "success", "validation_id": validation_id, "scope": scope, "decision": decision, "employees_marked_for_reassessment": affected}


@app.post("/api/assessment-validations/{validation_id}/revoke")
def revoke_assessment_validation(validation_id: int, payload: RevokeDecision) -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        row = require_row(con, "assessment_validations", "validation_id", validation_id, "Assessment Validation")
        if not row["active_flag"]:
            return {"status": "success", "validation_id": validation_id, "already_revoked": True}
        con.execute(
            "UPDATE assessment_validations SET active_flag=0,revoked_by=?,revoked_at=CURRENT_TIMESTAMP WHERE validation_id=?",
            (user["user_id"], validation_id),
        )
        if row["scope"] == "Employee" and row["employee_pk"]:
            mark_employee_for_reassessment(con, int(row["employee_pk"]), user["user_id"])
            affected = 1
        else:
            req = require_row(con, "training_requirements", "requirement_id", int(row["requirement_id"]), "Training Requirement")
            affected = mark_position_for_reassessment(con, int(req["position_standard_id"] or req["position_id"]), user["user_id"])
        con.commit()
    write_audit(user, "REVOKE_MAPPING_VALIDATION", "Assessment Validation", str(validation_id), payload.remarks or "")
    return {"status": "success", "validation_id": validation_id, "employees_marked_for_reassessment": affected}


@app.get("/api/assessment-overrides")
def list_assessment_overrides(active_only: bool = True, limit: int = 1000) -> list[dict[str, Any]]:
    where = "WHERE ao.active_flag=1" if active_only else ""
    df = query_df(
        f"""
        SELECT ao.*, e.employee_name, tc.training_name
        FROM assessment_overrides ao
        JOIN employees e ON e.employee_pk=ao.employee_pk
        JOIN training_requirements tr ON tr.requirement_id=ao.requirement_id
        JOIN training_catalog tc ON tc.training_id=tr.training_id
        {where}
        ORDER BY ao.approved_at DESC, ao.override_id DESC LIMIT ?
        """,
        (min(max(limit, 1), 5000),),
    ).fillna("")
    return df.to_dict(orient="records")


@app.post("/api/assessment-overrides")
def create_assessment_override(payload: AssessmentOverrideCreate) -> dict[str, Any]:
    user = current_user()
    allowed = {"Covered", "Potential Gap", "Validation Required", "Not Assessed", "Expired"}
    if payload.override_status not in allowed:
        raise HTTPException(status_code=400, detail=f"override_status harus salah satu: {', '.join(sorted(allowed))}")
    if not payload.reason.strip():
        raise HTTPException(status_code=400, detail="Reason wajib diisi untuk override.")
    with sqlite3.connect(DB_PATH) as con:
        require_row(con, "employees", "employee_pk", payload.employee_pk, "Employee")
        require_row(con, "training_requirements", "requirement_id", payload.requirement_id, "Training Requirement")
        con.execute(
            "UPDATE assessment_overrides SET active_flag=0,revoked_by=?,revoked_at=CURRENT_TIMESTAMP WHERE employee_pk=? AND requirement_id=? AND active_flag=1",
            (user["user_id"], payload.employee_pk, payload.requirement_id),
        )
        cur = con.execute(
            """INSERT INTO assessment_overrides(
                employee_pk,requirement_id,override_status,reason,evidence_reference,active_flag,approved_by,approved_at
            ) VALUES (?,?,?,?,?,1,?,CURRENT_TIMESTAMP)""",
            (payload.employee_pk, payload.requirement_id, payload.override_status, payload.reason.strip(), payload.evidence_reference, user["user_id"]),
        )
        override_id = int(cur.lastrowid)
        mark_employee_for_reassessment(con, payload.employee_pk, user["user_id"])
        con.commit()
    write_audit(user, "CREATE_ASSESSMENT_OVERRIDE", "Assessment Override", str(override_id), f"employee={payload.employee_pk}; requirement={payload.requirement_id}; status={payload.override_status}; reason={payload.reason}")
    return {"status": "success", "override_id": override_id, "assessment_state": "Pending Reassessment"}


@app.post("/api/assessment-overrides/{override_id}/revoke")
def revoke_assessment_override(override_id: int, payload: RevokeDecision) -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        row = require_row(con, "assessment_overrides", "override_id", override_id, "Assessment Override")
        if not row["active_flag"]:
            return {"status": "success", "override_id": override_id, "already_revoked": True}
        con.execute(
            "UPDATE assessment_overrides SET active_flag=0,revoked_by=?,revoked_at=CURRENT_TIMESTAMP WHERE override_id=?",
            (user["user_id"], override_id),
        )
        mark_employee_for_reassessment(con, int(row["employee_pk"]), user["user_id"])
        con.commit()
    write_audit(user, "REVOKE_ASSESSMENT_OVERRIDE", "Assessment Override", str(override_id), payload.remarks or "")
    return {"status": "success", "override_id": override_id, "assessment_state": "Pending Reassessment"}


@app.get("/api/competency/{employee_pk}/requirement/{requirement_id}/decision-trace")
def competency_decision_trace(employee_pk: int, requirement_id: int) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        row = con.execute(
            """
            SELECT da.*, ca.expiry_status, ca.certification_candidate
            FROM decision_assessment_current da
            LEFT JOIN competency_assessment ca
              ON ca.employee_pk=da.employee_pk AND ca.requirement_id=da.requirement_id
            WHERE da.employee_pk=? AND da.requirement_id=?
            ORDER BY da.decision_id DESC LIMIT 1
            """,
            (employee_pk, requirement_id),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Decision trace belum tersedia. Jalankan assessment/reassessment dengan Decision Engine v2 terlebih dahulu.")
        data = dict(row)
        try:
            data["evidence"] = __import__("json").loads(data.get("evidence_summary") or "{}")
        except Exception:
            data["evidence"] = {"raw": data.get("evidence_summary")}
        validation = con.execute("SELECT * FROM assessment_validations WHERE validation_id=?", (data.get("mapping_validation_id"),)).fetchone() if data.get("mapping_validation_id") else None
        override = con.execute("SELECT * FROM assessment_overrides WHERE override_id=?", (data.get("override_id"),)).fetchone() if data.get("override_id") else None
        data["validation"] = dict(validation) if validation else None
        data["override"] = dict(override) if override else None
        return data


@app.get("/api/validation-queue")
def validation_queue(limit: int = 1000) -> list[dict[str, Any]]:
    df = query_df(
        """
        SELECT ca.assessment_id,ca.employee_pk,ca.requirement_id,e.employee_name,
               ca.position_tna AS position_name,ca.training_name,
               ca.coverage_status AS status,ca.priority,ca.recommendation,
               ca.mapping_category,ca.certification_candidate
        FROM competency_assessment ca
        JOIN employees e ON e.employee_pk=ca.employee_pk
        WHERE ca.coverage_status='Validation Required' AND e.active_flag=1
        ORDER BY e.employee_name,ca.training_name
        LIMIT ?
        """, (min(max(limit,1),5000),)
    ).fillna("")
    return df.to_dict(orient="records")

@app.get("/api/definitions")
def definitions() -> dict[str, Any]:
    return {
        "Dashboard": [
            {"term":"Employees","definition":"Jumlah karyawan aktif pada Employee Master."},
            {"term":"Certifications","definition":"Jumlah seluruh record sertifikasi yang tersimpan."},
            {"term":"Expired","definition":"Sertifikasi dengan expiry date yang sudah melewati tanggal analisis."},
            {"term":"Near Expiry ≤90d","definition":"Sertifikasi yang akan berakhir dalam maksimal 90 hari."},
            {"term":"Training History","definition":"Record training aktual yang sudah tercatat pada sistem."},
        ],
        "Training": [
            {"term":"Training Requirement","definition":"Pelatihan yang dipersyaratkan untuk suatu jabatan."},
            {"term":"Employee","definition":"Karyawan aktif yang terpetakan pada posisi terkait."},
            {"term":"Mandatory","definition":"Training yang ditetapkan sebagai kewajiban."},
            {"term":"Additional","definition":"Training tambahan yang direkomendasikan sesuai kebutuhan."},
        ],
        "Certification": [
            {"term":"Active","definition":"Sertifikasi masih berlaku pada tanggal analisis."},
            {"term":"Near Expiry","definition":"Sertifikasi akan berakhir dalam ≤90 hari."},
            {"term":"Expired","definition":"Sertifikasi telah melewati expiry date."},
            {"term":"In Process","definition":"Sertifikasi masih tercatat dalam proses."},
        ],
        "Competency": [
            {"term":"Covered","definition":"Requirement memiliki evidence pemenuhan yang sesuai pada data sistem."},
            {"term":"Potential Gap","definition":"Indikasi belum terpenuhinya requirement berdasarkan evidence yang tersedia; bukan gap final."},
            {"term":"Validation Required","definition":"Hubungan evidence atau mapping belum cukup pasti sehingga memerlukan validasi TCD sebelum kesimpulan fulfillment dibuat."},
            {"term":"Assessment Source","definition":"System, Validated, Override, atau Legacy; menunjukkan bagaimana hasil assessment dibentuk."},
            {"term":"Not Assessed","definition":"Data atau mapping belum cukup untuk menghasilkan assessment."},
        ],
        "Assessment Queue": [
            {"term":"Pending Initial Assessment","definition":"Karyawan baru; data awal belum dikonfirmasi sehingga sistem belum menghasilkan coverage/gap."},
            {"term":"Assessment Ready","definition":"Data awal sudah dikonfirmasi dan Initial Assessment siap dijalankan."},
            {"term":"Pending Reassessment","definition":"Ada perubahan evidence, position, atau requirement setelah assessment sebelumnya dan perlu dihitung ulang."},
            {"term":"Assessed","definition":"Decision Engine telah menjalankan assessment terakhir untuk employee."},
        ],
        "Training History": [
            {"term":"Training History","definition":"Riwayat pelatihan aktual yang telah dilaksanakan dan dicatat."},
            {"term":"Completion Date","definition":"Tanggal penyelesaian training yang tercatat."},
            {"term":"Provider","definition":"Penyedia atau pelaksana training bila tersedia."},
        ],
        "AI Assistant": [
            {"term":"Conversation Context","definition":"Pertanyaan lanjutan mempertahankan employee dan objek yang sudah terverifikasi sampai percakapan direset."},
            {"term":"Data Source","definition":"Requirement, Training History, fulfillment, certification, competency, dan Action Center dipisahkan sesuai maksud pertanyaan."},
            {"term":"Evidence Guard","definition":"Jawaban hanya menggunakan record yang tersedia; hasil kosong tidak diganti dengan data dari sumber lain."},
        ],
        "Data Management": [
            {"term":"Safe Bulk Import","definition":"Upload Excel divalidasi dan dipreview terlebih dahulu; database belum berubah sebelum Confirm Import."},
            {"term":"Template","definition":"Format Excel resmi per dataset agar nama kolom, tanggal, dan status dapat dibaca konsisten oleh sistem."},
            {"term":"Duplicate","definition":"Record yang kemungkinan sudah ada. Default aman adalah Skip; Update Existing hanya dilakukan jika dipilih saat konfirmasi."},
            {"term":"Import Batch","definition":"Satu proses validasi/import yang memiliki Batch ID untuk traceability dan audit."},
        ],
    }

@app.get("/api/data-management/logs")
def data_logs(limit: int = 100) -> list[dict[str, Any]]:
    df = query_df(
        "SELECT source_name AS source_file, source_type AS dataset_type, record_count AS rows_read, imported_at, notes AS message FROM data_import_log ORDER BY imported_at DESC LIMIT ?",
        (min(max(limit,1),500),)
    ).fillna("")
    # Keep a stable UI contract for imported/rejected counts when the schema stores them in notes.
    def parse_note(s):
        import re
        m1=re.search(r'imported=(\d+)', str(s)); m2=re.search(r'rejected=(\d+)', str(s)); m3=re.search(r'status=([^;]+)', str(s))
        return (int(m1.group(1)) if m1 else 0, int(m2.group(1)) if m2 else 0, m3.group(1).strip() if m3 else '')
    parsed=df['message'].apply(parse_note) if not df.empty else pd.Series([],dtype=object)
    if not df.empty:
        df['rows_imported']=[x[0] for x in parsed]; df['rows_rejected']=[x[1] for x in parsed]; df['status']=[x[2] for x in parsed]
    return df.to_dict(orient='records')

@app.get("/api/data-management/templates")
def import_templates() -> list[dict[str, Any]]:
    from bulk_import import DATASET_CONFIG
    result=[]
    for dataset,cfg in DATASET_CONFIG.items():
        path=TEMPLATE_DIR / cfg["template"]
        result.append({"dataset_type":dataset,"filename":cfg["template"],"available":path.exists()})
    return result

@app.get("/api/data-management/templates/{dataset_type}")
def download_import_template(dataset_type: str):
    from bulk_import import DATASET_CONFIG
    if dataset_type not in DATASET_CONFIG:
        raise HTTPException(status_code=404, detail="Template dataset tidak ditemukan.")
    path=TEMPLATE_DIR / DATASET_CONFIG[dataset_type]["template"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="File template belum tersedia.")
    return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=path.name)

@app.post("/api/data-management/import/validate")
async def validate_bulk_import(dataset_type: str = Query(...), file: UploadFile = File(...)) -> dict[str, Any]:
    from bulk_import import DATASET_CONFIG, read_template_excel, create_validation_batch
    if dataset_type not in DATASET_CONFIG:
        raise HTTPException(status_code=400, detail="Dataset type tidak didukung.")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Gunakan file Excel .xlsx dari template resmi.")
    raw=await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="File kosong.")
    user=current_user()
    try:
        df=read_template_excel(raw,dataset_type)
        if df.empty:
            raise ValueError("Tidak ada baris data pada sheet DATA_ENTRY.")
        if len(df)>5000:
            raise ValueError("Maksimum 5.000 baris per batch untuk build v22.10.5.")
        with sqlite3.connect(DB_PATH) as con:
            summary=create_validation_batch(con,dataset_type,file.filename or "upload.xlsx",df,user["user_id"])
        write_audit(user,"VALIDATE_IMPORT",dataset_type,summary["batch_id"],f"file={file.filename}; rows={len(df)}; valid={summary['rows_valid']}; warning={summary['rows_warning']}; duplicate={summary['rows_duplicate']}; rejected={summary['rows_rejected']}")
        return summary
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))

@app.get("/api/data-management/import/batches")
def list_import_batches(limit: int = 50) -> list[dict[str, Any]]:
    from bulk_import import recent_batches
    with sqlite3.connect(DB_PATH) as con:
        return recent_batches(con,limit)

@app.get("/api/data-management/import/batches/{batch_id}")
def get_import_batch(batch_id: str) -> dict[str, Any]:
    from bulk_import import batch_summary
    try:
        with sqlite3.connect(DB_PATH) as con:
            return batch_summary(con,batch_id,include_rows=True)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc))

@app.post("/api/data-management/import/batches/{batch_id}/confirm")
def confirm_bulk_import(batch_id: str, payload: ImportConfirm) -> dict[str, Any]:
    from bulk_import import confirm_batch
    if payload.duplicate_action not in {"skip","update"}:
        raise HTTPException(status_code=400, detail="duplicate_action harus skip atau update.")
    user=current_user()
    try:
        safety=create_backup(f"pre_import_{batch_id}",user["user_id"])
        with sqlite3.connect(DB_PATH) as con:
            result=confirm_batch(con,batch_id,user["user_id"],payload.accept_warnings,payload.duplicate_action)
        result["safety_backup"]=safety["filename"]
        write_audit(user,"CONFIRM_IMPORT",result["dataset_type"],batch_id,f"imported={result['rows_imported']}; skipped={result['rows_skipped']}; affected_employees={result['affected_employees']}")
        return result
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))

@app.get("/api/data-management/import/batches/{batch_id}/report")
def import_batch_report(batch_id: str):
    from bulk_import import batch_summary
    try:
        with sqlite3.connect(DB_PATH) as con:
            data=batch_summary(con,batch_id,include_rows=True,limit=5000)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc))
    rows=[]
    for r in data.get("rows",[]):
        base={"row_number":r["row_number"],"status":r["row_status"],"proposed_action":r.get("proposed_action",""),"issues":" | ".join(i.get("message","") for i in r.get("issues",[]))}
        base.update(r.get("row_data",{})); rows.append(base)
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer,index=False,sheet_name="VALIDATION_RESULT")
        pd.DataFrame([{k:v for k,v in data.items() if k!="rows"}]).to_excel(writer,index=False,sheet_name="BATCH_INFO")
    out.seek(0)
    headers={"Content-Disposition":f'attachment; filename="{batch_id}_validation_report.xlsx"'}
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers=headers)

@app.post("/api/data-management/import", deprecated=True)
async def data_import_legacy(dataset_type: str = Query(...), file: UploadFile = File(...)) -> dict[str, Any]:
    raise HTTPException(status_code=410, detail="Direct import dinonaktifkan pada v22.10.5. Gunakan /api/data-management/import/validate lalu Confirm Import agar data melalui validation, preview, duplicate check, dan audit batch.")

@app.post("/api/ai/query")
def ai_query(payload: AIQuery) -> dict[str, Any]:
    try:
        user = current_user()
        conversation_id = (payload.conversation_id or secrets.token_urlsafe(12))[:80]
        context_key = f"{user['user_id']}:{conversation_id}"
        reset_phrases = {"hapus konteks", "reset konteks", "percakapan baru", "new conversation"}
        if payload.reset_context or payload.question.strip().lower() in reset_phrases:
            _AI_CONTEXTS.pop(context_key, None)
            return {
                "question": payload.question, "user": user,
                "response": "Konteks percakapan sudah dihapus. Silakan mulai pertanyaan baru.",
                "status": "CONTEXT_RESET", "intents": [], "intent_details": [], "agents": [],
                "result_count": 0, "evidence_total": 0, "unique_employee_count": 0,
                "results": {}, "decision_summary": [], "decision_trace": [],
                "conversation_id": conversation_id,
            }
        if is_management_analytics_question(payload.question):
            with sqlite3.connect(DB_PATH) as con:
                data = management_analytics(con, refresh_snapshot=True, created_by=user["user_id"])
                con.commit()
            write_audit(user, "AI_ANALYTICS_QUERY", target="Management Analytics", details=payload.question[:500])
            return {
                "question": payload.question, "user": user, "response": management_ai_summary(data),
                "status": "ANALYTICS", "intents": ["management_analytics"], "agents": ["Analytics Engine"],
                "result_count": 0, "evidence_total": 0, "unique_employee_count": data["summary"].get("active_employees",0),
                "results": {}, "decision_summary": [], "decision_trace": [], "analytics": data["summary"],
                "data_readiness": data["data_readiness"],
                "conversation_id": conversation_id,
            }
        from ai_assistant import run_query, response_text
        result = run_query(DB_PATH, payload.question, context=_AI_CONTEXTS.get(context_key, {}))
        _AI_CONTEXTS[context_key] = dict(result.get("conversation_context", {}))
        if len(_AI_CONTEXTS) > 500:
            for old_key in list(_AI_CONTEXTS)[:100]:
                _AI_CONTEXTS.pop(old_key, None)
        write_audit(user, "AI_QUERY", target="AI Assistant", details=payload.question[:500])
        return {
            "question": payload.question,
            "user": user,
            "response": response_text(payload.question, result),
            "status": result.get("status"),
            "intents": result.get("intents", []),
            "intent_details": result.get("intent_details", []),
            "agents": result.get("agents", []),
            "result_count": result.get("total", 0),
            "evidence_total": result.get("evidence_total", 0),
            "unique_employee_count": result.get("unique_employee_count", 0),
            "results": {
                agent: {
                    "mode": data["mode"],
                    "records": data["data"].to_dict(orient="records"),
                }
                for agent, data in result.get("results", {}).items()
            },
            "decision_summary": result.get("decision_summary", []),
            "decision_trace": result.get("decision_trace", []),
            "conversation_id": conversation_id,
            "source_domain": result.get("conversation_context", {}).get("domain", ""),
        }
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))

@app.post("/api/ai/export")
def ai_export(payload: AIQuery):
    try:
        user = current_user()
        if is_management_analytics_question(payload.question):
            with sqlite3.connect(DB_PATH) as con:
                data = management_analytics(con, refresh_snapshot=True, created_by=user["user_id"])
                con.commit()
            frames = analytics_frames(data)
            buf = io.BytesIO()
            with pd.ExcelWriter(buf, engine="openpyxl") as writer:
                pd.DataFrame([{"Question": payload.question, "Generated At": data["generated_at"], "Build": BUILD}]).to_excel(writer, sheet_name="AI Analytics Query", index=False)
                for sheet_name, frame in frames.items():
                    (frame if not frame.empty else pd.DataFrame([{"Information":"No records available."}])).to_excel(writer, sheet_name=sheet_name[:31], index=False)
                _style_management_workbook(writer)
            buf.seek(0)
            write_audit(user, "AI_ANALYTICS_EXPORT", "Reporting", "ai_management_analytics.xlsx", payload.question[:500])
            return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":"attachment; filename=ai_management_analytics.xlsx"})
        from ai_assistant import run_query
        export_context = {}
        if payload.conversation_id:
            export_context = _AI_CONTEXTS.get(f"{user['user_id']}:{payload.conversation_id}", {})
        result = run_query(DB_PATH, payload.question, context=export_context)
        decision = pd.DataFrame(result.get("decision_summary", []))
        evidence=[]
        for agent, data in result.get("results", {}).items():
            frame=data.get("data")
            if frame is not None and not frame.empty:
                f=frame.copy(); f.insert(0,"source",data.get("mode") or agent.replace(" Agent","")); evidence.append(f)
        trace=pd.DataFrame(result.get("decision_trace", []))
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            if not decision.empty:
                decision.to_excel(writer, sheet_name="Decision Summary", index=False)
            if evidence:
                pd.concat(evidence, ignore_index=True, sort=False).to_excel(writer, sheet_name="Evidence", index=False)
            else:
                pd.DataFrame([{ "message":"Tidak ada evidence untuk diekspor.", "question":payload.question }]).to_excel(writer, sheet_name="Evidence", index=False)
            for agent, data in result.get("results", {}).items():
                frame=data.get("data")
                if frame is not None and not frame.empty:
                    frame.to_excel(writer, sheet_name=agent.replace(" Agent", "")[:31], index=False)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":"attachment; filename=ai_assistant_result.xlsx"},
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))

# -----------------------------------------------------------------------------
# Build v22.6 — Validation Workflow + Action Management + Evidence + Notifications
# -----------------------------------------------------------------------------

@app.get("/api/reference/pic-users")
def reference_pic_users() -> list[dict[str, Any]]:
    return query_df("SELECT user_id,display_name,role FROM app_users WHERE active_flag=1 AND user_id<>'U999' ORDER BY display_name").fillna("").to_dict(orient="records")


@app.get("/api/actions/meta")
def actions_meta() -> dict[str, Any]:
    return {"statuses": ACTION_STATUSES, "types": ACTION_TYPES, "priorities": ACTION_PRIORITIES}


@app.get("/api/actions")
def list_actions(status: str = "", pic_user_id: str = "", priority: str = "", employee_pk: int | None = None,
                 deleted: str = "active", limit: int = 2000) -> list[dict[str, Any]]:
    clauses=[]; params=[]
    if deleted == "active": clauses.append("COALESCE(a.is_deleted,0)=0")
    elif deleted == "deleted": clauses.append("COALESCE(a.is_deleted,0)=1")
    elif deleted != "all": raise HTTPException(status_code=400,detail="deleted harus active, deleted, atau all.")
    if status: clauses.append("a.status=?"); params.append(status)
    if pic_user_id: clauses.append("a.pic_user_id=?"); params.append(pic_user_id)
    if priority: clauses.append("a.priority=?"); params.append(priority)
    if employee_pk is not None: clauses.append("a.employee_pk=?"); params.append(employee_pk)
    where="WHERE "+" AND ".join(clauses) if clauses else ""
    params.append(min(max(limit,1),5000))
    df=query_df(f"""
        SELECT a.*,e.employee_name,p.position_name,tc.training_name,u.display_name AS pic_name,
               CASE WHEN a.due_date IS NOT NULL AND a.status NOT IN ('Completed','Closed','Cancelled','No Action Required')
                    THEN CAST(julianday(date(a.due_date))-julianday(date('now')) AS INTEGER) END AS days_remaining,
               (SELECT COUNT(*) FROM evidence_documents ed WHERE ed.entity_type='action' AND ed.entity_id=a.action_id AND ed.active_flag=1) AS evidence_count
        FROM actions a JOIN employees e ON e.employee_pk=a.employee_pk
        LEFT JOIN positions p ON p.position_id=e.current_position_id
        LEFT JOIN training_requirements tr ON tr.requirement_id=a.requirement_id
        LEFT JOIN training_catalog tc ON tc.training_id=tr.training_id
        LEFT JOIN app_users u ON u.user_id=a.pic_user_id
        {where}
        ORDER BY CASE a.status WHEN 'Open' THEN 1 WHEN 'Planned' THEN 2 WHEN 'Scheduled' THEN 3 WHEN 'In Progress' THEN 4 WHEN 'Waiting Evidence' THEN 5 WHEN 'Waiting External Party' THEN 6 ELSE 9 END,
                 CASE a.priority WHEN 'High' THEN 1 WHEN 'Validation' THEN 2 WHEN 'Medium' THEN 3 WHEN 'Low' THEN 4 ELSE 5 END,
                 COALESCE(a.due_date,'9999-12-31'),a.action_id DESC LIMIT ?
    """,tuple(params)).fillna("")
    return df.to_dict(orient="records")


@app.get("/api/actions/summary")
def actions_summary() -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        rows=con.execute("SELECT status,COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 GROUP BY status").fetchall()
        by_status={r[0]:int(r[1]) for r in rows}
        overdue=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND status NOT IN ('Completed','Closed','Cancelled','No Action Required') AND due_date IS NOT NULL AND date(due_date)<date('now')").fetchone()[0])
        high=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND priority='High' AND status NOT IN ('Completed','Closed','Cancelled','No Action Required')").fetchone()[0])
        unassigned=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND pic_user_id IS NULL AND status NOT IN ('Completed','Closed','Cancelled','No Action Required')").fetchone()[0])
    return {"by_status":by_status,"overdue":overdue,"high_priority":high,"unassigned":unassigned,"total":sum(by_status.values())}


@app.post("/api/actions")
def create_action_endpoint(payload: ActionCreate) -> dict[str, Any]:
    user=current_user()
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.execute("PRAGMA foreign_keys=ON")
            action_id=workflow_create_action(con,employee_pk=payload.employee_pk,requirement_id=payload.requirement_id,
                action_type=payload.action_type,title=payload.title,priority=payload.priority,recommendation=payload.recommendation,
                pic_user_id=payload.pic_user_id,due_date=iso(payload.due_date),remarks=payload.remarks,
                source_decision_id=payload.source_decision_id,created_by=user["user_id"])
            con.commit()
        write_audit(user,"CREATE_ACTION","Action",str(action_id),f"employee={payload.employee_pk}; type={payload.action_type}; priority={payload.priority}; pic={payload.pic_user_id or 'Unassigned'}")
        return {"status":"success","action_id":action_id}
    except ValueError as exc:
        raise HTTPException(status_code=400,detail=str(exc))


@app.put("/api/actions/{action_id}")
def update_action_endpoint(action_id: int,payload: ActionUpdate) -> dict[str, Any]:
    user=current_user()
    try:
        with sqlite3.connect(DB_PATH) as con:
            con.execute("PRAGMA foreign_keys=ON")
            result=workflow_update_action(con,action_id,status=payload.status,pic_user_id=payload.pic_user_id,
                due_date=iso(payload.due_date),priority=payload.priority,remarks=payload.remarks,
                no_action_reason=payload.no_action_reason,changed_by=user["user_id"])
            con.commit()
        write_audit(user,"UPDATE_ACTION","Action",str(action_id),f"status={result['new_status']}; pic={result.get('pic_user_id') or 'Unassigned'}")
        return {"status":"success",**result}
    except ValueError as exc:
        raise HTTPException(status_code=400,detail=str(exc))


@app.get("/api/actions/{action_id}")
def action_detail(action_id: int) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory=sqlite3.Row
        a=con.execute("""SELECT a.*,e.employee_name,p.position_name,tc.training_name,u.display_name pic_name
                         FROM actions a JOIN employees e ON e.employee_pk=a.employee_pk LEFT JOIN positions p ON p.position_id=e.current_position_id
                         LEFT JOIN training_requirements tr ON tr.requirement_id=a.requirement_id LEFT JOIN training_catalog tc ON tc.training_id=tr.training_id
                         LEFT JOIN app_users u ON u.user_id=a.pic_user_id
                         WHERE a.action_id=? AND COALESCE(a.is_deleted,0)=0""",(action_id,)).fetchone()
        if not a: raise HTTPException(status_code=404,detail="Action tidak ditemukan.")
        history=[dict(r) for r in con.execute("SELECT * FROM action_history WHERE action_id=? ORDER BY changed_at DESC,history_id DESC",(action_id,)).fetchall()]
        docs=[dict(r) for r in con.execute("""SELECT ed.document_id,ed.document_type,ed.original_filename,ed.mime_type,ed.file_size,
                                      ed.description,ed.uploaded_by,COALESCE(u.display_name,ed.uploaded_by) uploaded_by_name,ed.uploaded_at,
                                      CASE WHEN lower(COALESCE(ed.mime_type,'')) LIKE 'image/%'
                                             OR lower(COALESCE(ed.mime_type,''))='application/pdf' THEN 1 ELSE 0 END viewable
                               FROM evidence_documents ed LEFT JOIN app_users u ON u.user_id=ed.uploaded_by
                               WHERE ed.entity_type='action' AND ed.entity_id=? AND ed.active_flag=1
                               ORDER BY ed.uploaded_at DESC,ed.document_id DESC""",(action_id,)).fetchall()]
        readiness=action_completion_ready(con,a)
        return {"action":dict(a),"history":history,"evidence":docs,"completion_readiness":readiness}


@app.delete("/api/actions/{action_id}")
def delete_action_endpoint(action_id: int, payload: ActionDelete) -> dict[str, Any]:
    user=current_user(); reason=payload.reason.strip()
    if not reason:
        raise HTTPException(status_code=400,detail="Deletion Reason wajib diisi.")
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory=sqlite3.Row
        action=con.execute("SELECT * FROM actions WHERE action_id=?",(action_id,)).fetchone()
        if not action: raise HTTPException(status_code=404,detail="Action tidak ditemukan.")
        if int(action["is_deleted"] or 0): raise HTTPException(status_code=409,detail="Action sudah dihapus.")
        con.execute("""UPDATE actions SET is_deleted=1,deleted_at=CURRENT_TIMESTAMP,deleted_by=?,deletion_reason=?,
                       restored_at=NULL,restored_by=NULL,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE action_id=?""",
                    (user["user_id"],reason,user["user_id"],action_id))
        con.execute("""INSERT INTO action_history(action_id,event_type,old_status,new_status,old_pic_user_id,new_pic_user_id,remarks,changed_by)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (action_id,"Soft Deleted",action["status"],action["status"],action["pic_user_id"],action["pic_user_id"],reason,user["user_id"]))
        con.execute("UPDATE notifications SET active_flag=0 WHERE source_type='action' AND source_id=?",(action_id,))
        con.commit()
    write_audit(user,"SOFT_DELETE_ACTION","Action",str(action_id),f"reason={reason}")
    return {"status":"success","action_id":action_id,"deleted":True,"evidence_retained":True}


@app.post("/api/actions/{action_id}/restore")
def restore_action_endpoint(action_id: int, payload: ActionRestore) -> dict[str, Any]:
    user=current_user(); reason=payload.reason.strip()
    if not reason:
        raise HTTPException(status_code=400,detail="Restore Reason wajib diisi.")
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory=sqlite3.Row
        action=con.execute("SELECT * FROM actions WHERE action_id=?",(action_id,)).fetchone()
        if not action: raise HTTPException(status_code=404,detail="Action tidak ditemukan.")
        if not int(action["is_deleted"] or 0): raise HTTPException(status_code=409,detail="Action masih aktif.")
        con.execute("""UPDATE actions SET is_deleted=0,restored_at=CURRENT_TIMESTAMP,restored_by=?,
                       deleted_at=NULL,deleted_by=NULL,deletion_reason=NULL,updated_by=?,updated_at=CURRENT_TIMESTAMP
                       WHERE action_id=?""",(user["user_id"],user["user_id"],action_id))
        con.execute("""INSERT INTO action_history(action_id,event_type,old_status,new_status,old_pic_user_id,new_pic_user_id,remarks,changed_by)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (action_id,"Restored",action["status"],action["status"],action["pic_user_id"],action["pic_user_id"],reason,user["user_id"]))
        refresh_notifications(con)
        con.commit()
    write_audit(user,"RESTORE_ACTION","Action",str(action_id),f"reason={reason}")
    return {"status":"success","action_id":action_id,"restored":True}


@app.post("/api/evidence/upload")
async def upload_evidence(entity_type: str=Form(...),entity_id: int=Form(...),document_type: str=Form("Other"),description: str=Form(""),file: UploadFile=File(...)) -> dict[str, Any]:
    user=current_user(); et=entity_type.strip().lower()
    try:
        suffix=Path(file.filename or "").suffix.lower()
        if suffix not in ALLOWED_EVIDENCE_EXT:
            raise ValueError("Tipe file evidence tidak didukung. Gunakan PDF, PNG/JPG, Excel/CSV, atau DOCX.")
        content=await file.read()
        if not content: raise ValueError("File evidence kosong.")
        if len(content)>MAX_EVIDENCE_BYTES: raise ValueError("Ukuran evidence maksimum 15 MB per file.")
        with sqlite3.connect(DB_PATH) as con:
            con.execute("PRAGMA foreign_keys=ON")
            assert_entity(con,et,entity_id)
            employee_pk,requirement_id=infer_entity_context(con,et,entity_id)
            stored=_safe_filename(file.filename or "evidence")
            folder=ROOT/"evidence"/(et if et in {"action","validation","training","certification"} else "other")
            folder.mkdir(parents=True,exist_ok=True)
            path=folder/stored; path.write_bytes(content)
            rel=str(path.relative_to(ROOT)).replace("\\","/")
            cur=con.execute("""INSERT INTO evidence_documents(entity_type,entity_id,employee_pk,requirement_id,document_type,original_filename,stored_filename,file_path,mime_type,file_size,description,uploaded_by)
                               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",(et,entity_id,employee_pk,requirement_id,document_type,file.filename or stored,stored,rel,file.content_type,len(content),description or None,user["user_id"]))
            doc_id=int(cur.lastrowid); con.commit()
        write_audit(user,"UPLOAD_EVIDENCE","Evidence",str(doc_id),f"entity={et}:{entity_id}; file={file.filename}; type={document_type}")
        return {"status":"success","document_id":doc_id,"entity_type":et,"entity_id":entity_id,"file_name":file.filename,"file_size":len(content)}
    except ValueError as exc:
        raise HTTPException(status_code=400,detail=str(exc))


@app.get("/api/evidence")
def list_evidence(entity_type: str="",entity_id: int | None=None,employee_pk: int | None=None,limit: int=1000) -> list[dict[str, Any]]:
    clauses=["active_flag=1", "(entity_type<>'action' OR EXISTS (SELECT 1 FROM actions a WHERE a.action_id=evidence_documents.entity_id AND COALESCE(a.is_deleted,0)=0))"]; params=[]
    if entity_type: clauses.append("entity_type=?");params.append(entity_type.lower())
    if entity_id is not None: clauses.append("entity_id=?");params.append(entity_id)
    if employee_pk is not None: clauses.append("employee_pk=?");params.append(employee_pk)
    params.append(min(max(limit,1),5000))
    return query_df(f"SELECT document_id,entity_type,entity_id,employee_pk,requirement_id,document_type,original_filename,mime_type,file_size,description,uploaded_by,uploaded_at FROM evidence_documents WHERE {' AND '.join(clauses)} ORDER BY uploaded_at DESC LIMIT ?",tuple(params)).fillna("").to_dict(orient="records")


@app.get("/api/evidence/{document_id}/download")
def download_evidence(document_id: int):
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory=sqlite3.Row
        row=con.execute("SELECT * FROM evidence_documents WHERE document_id=? AND active_flag=1",(document_id,)).fetchone()
        if not row: raise HTTPException(status_code=404,detail="Evidence tidak ditemukan.")
        if str(row["entity_type"] or "").lower()=="action" and not con.execute("SELECT 1 FROM actions WHERE action_id=? AND COALESCE(is_deleted,0)=0",(row["entity_id"],)).fetchone():
            raise HTTPException(status_code=404,detail="Evidence tidak tersedia karena action telah dihapus.")
        path=(ROOT/row["file_path"]).resolve()
        evidence_root=(ROOT/"evidence").resolve()
        if evidence_root not in path.parents or not path.exists(): raise HTTPException(status_code=404,detail="File evidence tidak ditemukan pada storage.")
        return FileResponse(path,filename=row["original_filename"],media_type=row["mime_type"] or "application/octet-stream")


@app.get("/api/evidence/{document_id}/view")
def view_evidence(document_id: int):
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory=sqlite3.Row
        row=con.execute("SELECT * FROM evidence_documents WHERE document_id=? AND active_flag=1",(document_id,)).fetchone()
        if not row: raise HTTPException(status_code=404,detail="Evidence tidak ditemukan.")
        if str(row["entity_type"] or "").lower()=="action" and not con.execute("SELECT 1 FROM actions WHERE action_id=? AND COALESCE(is_deleted,0)=0",(row["entity_id"],)).fetchone():
            raise HTTPException(status_code=404,detail="Evidence tidak tersedia karena action telah dihapus.")
        path=(ROOT/row["file_path"]).resolve()
        evidence_root=(ROOT/"evidence").resolve()
        if evidence_root not in path.parents or not path.exists(): raise HTTPException(status_code=404,detail="File evidence tidak ditemukan pada storage.")
        mime=(row["mime_type"] or "application/octet-stream").lower()
        if not (mime.startswith("image/") or mime=="application/pdf"):
            raise HTTPException(status_code=415,detail="Preview hanya tersedia untuk PDF dan image. Gunakan Download untuk file ini.")
        safe_name=str(row["original_filename"] or "evidence").replace('"','')
        return FileResponse(path,media_type=mime,headers={"Content-Disposition":f'inline; filename="{safe_name}"'})


@app.post("/api/evidence/{document_id}/archive")
def archive_evidence(document_id: int) -> dict[str, Any]:
    user=current_user()
    with sqlite3.connect(DB_PATH) as con:
        if not con.execute("SELECT 1 FROM evidence_documents WHERE document_id=?",(document_id,)).fetchone(): raise HTTPException(status_code=404,detail="Evidence tidak ditemukan.")
        con.execute("UPDATE evidence_documents SET active_flag=0 WHERE document_id=?",(document_id,));con.commit()
    write_audit(user,"ARCHIVE_EVIDENCE","Evidence",str(document_id),"")
    return {"status":"success","document_id":document_id}


@app.post("/api/notifications/refresh")
def notification_refresh() -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        result=refresh_notifications(con);con.commit()
    return {"status":"success",**result}


@app.get("/api/notifications")
def list_notifications(unread_only: bool=False,include_dismissed: bool=False,limit: int=300) -> list[dict[str, Any]]:
    clauses=["n.active_flag=1"];params=[]
    if unread_only: clauses.append("n.read_flag=0")
    if not include_dismissed: clauses.append("n.dismissed_at IS NULL")
    params.append(min(max(limit,1),1000))
    df=query_df(f"""SELECT n.*,e.employee_name FROM notifications n LEFT JOIN employees e ON e.employee_pk=n.employee_pk
                      WHERE {' AND '.join(clauses)} ORDER BY CASE n.severity WHEN 'Critical' THEN 1 WHEN 'High' THEN 2 WHEN 'Attention' THEN 3 ELSE 4 END,
                      COALESCE(n.due_date,'9999-12-31'),n.created_at DESC LIMIT ?""",tuple(params)).fillna("")
    return df.to_dict(orient="records")


@app.post("/api/notifications/{notification_id}/read")
def notification_read(notification_id: int,payload: NotificationRead) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        if not con.execute("SELECT 1 FROM notifications WHERE notification_id=?",(notification_id,)).fetchone(): raise HTTPException(status_code=404,detail="Notification tidak ditemukan.")
        con.execute("UPDATE notifications SET read_flag=?,read_at=CASE WHEN ?=1 THEN CURRENT_TIMESTAMP ELSE NULL END,updated_at=CURRENT_TIMESTAMP WHERE notification_id=?",(1 if payload.read else 0,1 if payload.read else 0,notification_id));con.commit()
    return {"status":"success","notification_id":notification_id,"read":payload.read}


@app.post("/api/notifications/{notification_id}/dismiss")
def notification_dismiss(notification_id: int) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        if not con.execute("SELECT 1 FROM notifications WHERE notification_id=?",(notification_id,)).fetchone(): raise HTTPException(status_code=404,detail="Notification tidak ditemukan.")
        con.execute("UPDATE notifications SET dismissed_at=CURRENT_TIMESTAMP,read_flag=1,read_at=COALESCE(read_at,CURRENT_TIMESTAMP),updated_at=CURRENT_TIMESTAMP WHERE notification_id=?",(notification_id,));con.commit()
    return {"status":"success","notification_id":notification_id}



# Build v22.10.5 — safe test-employee purge hotfix
@app.get("/api/employees/{employee_pk}/profile")
def employee_profile(employee_pk: int) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row
        emp = con.execute("""SELECT e.*,p.position_name,COALESCE(e.department,p.department,'') department
                             FROM employees e LEFT JOIN positions p ON p.position_id=e.current_position_id
                             WHERE e.employee_pk=?""", (employee_pk,)).fetchone()
        if not emp:
            raise HTTPException(status_code=404, detail="Employee tidak ditemukan.")
        position_id = int(emp["current_position_id"] or 0)
        requirement_count = int(con.execute("""SELECT COUNT(*) FROM training_requirements
                    WHERE active_flag=1 AND COALESCE(requirement_status,'Active')='Active'
                    AND COALESCE(position_standard_id,position_id)=?""", (position_id,)).fetchone()[0]) if position_id else 0
        training_count = int(con.execute("SELECT COUNT(*) FROM training_history WHERE employee_pk=? AND COALESCE(record_status,'Active')<>'Archived'", (employee_pk,)).fetchone()[0])
        cert_count = int(con.execute("SELECT COUNT(*) FROM certifications WHERE employee_pk=? AND COALESCE(record_status,'Active')<>'Archived'", (employee_pk,)).fetchone()[0])
        open_actions = int(con.execute("SELECT COUNT(*) FROM actions WHERE employee_pk=? AND COALESCE(is_deleted,0)=0 AND status NOT IN ('Completed','Closed','Cancelled','No Action Required')", (employee_pk,)).fetchone()[0])
        overdue = int(con.execute("SELECT COUNT(*) FROM actions WHERE employee_pk=? AND COALESCE(is_deleted,0)=0 AND status NOT IN ('Completed','Closed','Cancelled','No Action Required') AND due_date IS NOT NULL AND date(due_date)<date('now')", (employee_pk,)).fetchone()[0])

        assessment_rows = pd.read_sql_query("""SELECT da.*,tr.requirement_type AS requirement_type_master,tc.training_name AS requirement_training_name
                  FROM decision_assessment_current da
                  LEFT JOIN training_requirements tr ON tr.requirement_id=da.requirement_id
                  LEFT JOIN training_catalog tc ON tc.training_id=tr.training_id
                  WHERE da.employee_pk=? ORDER BY CASE da.final_status WHEN 'Potential Gap' THEN 1 WHEN 'Expired' THEN 2 WHEN 'Validation Required' THEN 3 WHEN 'Not Assessed' THEN 4 ELSE 5 END, da.training_name""", con, params=(employee_pk,)).fillna("")
        if assessment_rows.empty:
            assessment_rows = pd.read_sql_query("""SELECT ca.*,ca.coverage_status final_status,ca.assessment_source,'' evidence_summary,'' engine_version
                    FROM competency_assessment ca WHERE ca.employee_pk=? ORDER BY ca.training_name""", con, params=(employee_pk,)).fillna("")
        distribution = {}
        if not assessment_rows.empty:
            col = 'final_status' if 'final_status' in assessment_rows.columns else 'coverage_status'
            distribution = {str(k): int(v) for k,v in assessment_rows[col].value_counts().to_dict().items()}

        requirements = pd.read_sql_query("""SELECT tr.requirement_id,tc.training_name,tr.requirement_type,tr.requirement_code,tr.delivery_type,tr.regulatory_flag,
                    COALESCE(tr.requirement_status,'Active') requirement_status,da.final_status coverage_status,da.priority,da.assessment_source,
                    da.evidence_summary,da.rationale
                FROM training_requirements tr JOIN training_catalog tc ON tc.training_id=tr.training_id
                LEFT JOIN decision_assessment_current da ON da.employee_pk=? AND da.requirement_id=tr.requirement_id
                WHERE tr.active_flag=1 AND COALESCE(tr.requirement_status,'Active')='Active' AND COALESCE(tr.position_standard_id,tr.position_id)=?
                ORDER BY CASE COALESCE(da.final_status,'') WHEN 'Potential Gap' THEN 1 WHEN 'Expired' THEN 2 WHEN 'Validation Required' THEN 3 WHEN 'Not Assessed' THEN 4 WHEN 'Covered' THEN 5 ELSE 6 END,tc.training_name""", con, params=(employee_pk,position_id)).fillna("").to_dict(orient="records") if position_id else []
        training = pd.read_sql_query("""SELECT th.history_id,COALESCE(th.training_name,th.training_name_raw,tc.training_name) training_name,
                    th.training_date,th.completion_date,COALESCE(th.result_status,th.completion_status) result_status,th.provider,th.valid_until,
                    th.certificate_reference,th.evidence_status,th.remarks,th.created_by,th.updated_by
                FROM training_history th LEFT JOIN training_catalog tc ON tc.training_id=COALESCE(th.training_id,th.training_pk)
                WHERE th.employee_pk=? AND COALESCE(th.record_status,'Active')<>'Archived' ORDER BY COALESCE(th.completion_date,th.training_date) DESC""", con, params=(employee_pk,)).fillna("").to_dict(orient="records")
        certifications = pd.read_sql_query("""SELECT certification_id,certification_name_raw,certificate_number,COALESCE(issuer,issuer_type) issuer,
                    COALESCE(issue_date,certification_date) issue_date,COALESCE(expiry_date,expired_date) expiry_date,
                    COALESCE(certification_status,status) certification_status,renewal_status,reason_code,remarks,previous_record_id,created_by,updated_by
                FROM certifications WHERE employee_pk=? AND COALESCE(record_status,'Active')<>'Archived'
                ORDER BY COALESCE(expiry_date,expired_date,'9999-12-31')""", con, params=(employee_pk,)).fillna("").to_dict(orient="records")
        actions = pd.read_sql_query("""SELECT a.*,tc.training_name,u.display_name pic_name,
                    CASE WHEN a.due_date IS NOT NULL AND a.status NOT IN ('Completed','Closed','Cancelled','No Action Required')
                         THEN CAST(julianday(date(a.due_date))-julianday(date('now')) AS INTEGER) END days_remaining
                FROM actions a LEFT JOIN training_requirements tr ON tr.requirement_id=a.requirement_id
                LEFT JOIN training_catalog tc ON tc.training_id=tr.training_id LEFT JOIN app_users u ON u.user_id=a.pic_user_id
                WHERE a.employee_pk=? AND COALESCE(a.is_deleted,0)=0 ORDER BY CASE WHEN a.status IN ('Closed','Cancelled','No Action Required') THEN 2 ELSE 1 END,COALESCE(a.due_date,'9999-12-31'),a.action_id DESC""", con, params=(employee_pk,)).fillna("").to_dict(orient="records")
        documents = pd.read_sql_query("""SELECT ed.document_id,ed.entity_type,ed.entity_id,ed.requirement_id,ed.document_type,ed.original_filename,ed.mime_type,
                    ed.description,ed.uploaded_by,ed.uploaded_at,COALESCE(u.display_name,ed.uploaded_by,'System') uploaded_by_name
                FROM evidence_documents ed LEFT JOIN app_users u ON u.user_id=ed.uploaded_by
                WHERE ed.employee_pk=? AND ed.active_flag=1
                  AND (ed.entity_type<>'action' OR EXISTS (SELECT 1 FROM actions ax WHERE ax.action_id=ed.entity_id AND COALESCE(ax.is_deleted,0)=0))
                ORDER BY ed.uploaded_at DESC""", con, params=(employee_pk,)).fillna("").to_dict(orient="records")
        position_history = pd.read_sql_query("""SELECT eph.employee_position_id,p.position_name,eph.effective_start,eph.effective_end,eph.is_current,eph.change_type,eph.change_reason,eph.remarks,eph.changed_by,eph.changed_at,
                    COALESCE(u.display_name,eph.changed_by,'System') changed_by_name
                FROM employee_position_history eph LEFT JOIN positions p ON p.position_id=eph.position_id LEFT JOIN app_users u ON u.user_id=eph.changed_by
                WHERE eph.employee_pk=? ORDER BY COALESCE(eph.effective_start,eph.changed_at) DESC""", con, params=(employee_pk,)).fillna("").to_dict(orient="records")

        # Enrich evidence links so Employee 360 shows the business object, not only a generic entity type.
        for doc in documents:
            et=str(doc.get("entity_type") or "").lower(); eid=int(doc.get("entity_id") or 0); req_id=int(doc.get("requirement_id") or 0)
            label=et.title() if et else "Evidence"; context=f"#{eid}" if eid else ""
            if et=="action" and eid:
                row=con.execute("SELECT title,status,requirement_id FROM actions WHERE action_id=?",(eid,)).fetchone()
                if row:
                    label=row[0] or "Action"; context=f"Action #{eid} · {row[1] or '—'}"; req_id=req_id or int(row[2] or 0)
            elif et=="training" and eid:
                row=con.execute("SELECT COALESCE(training_name,training_name_raw,'Training'),COALESCE(result_status,completion_status,'') FROM training_history WHERE history_id=?",(eid,)).fetchone()
                if row: label=row[0] or "Training"; context=f"Training History #{eid}" + (f" · {row[1]}" if row[1] else "")
            elif et=="certification" and eid:
                row=con.execute("SELECT COALESCE(certification_name_raw,'Certification'),COALESCE(certification_status,status,'') FROM certifications WHERE certification_id=?",(eid,)).fetchone()
                if row: label=row[0] or "Certification"; context=f"Certification #{eid}" + (f" · {row[1]}" if row[1] else "")
            elif et=="validation" and eid:
                row=con.execute("SELECT requirement_id,decision FROM assessment_validations WHERE validation_id=?",(eid,)).fetchone()
                if row: req_id=req_id or int(row[0] or 0); context=f"Validation #{eid}" + (f" · {row[1]}" if row[1] else "")
            if req_id:
                rr=con.execute("SELECT tc.training_name FROM training_requirements tr JOIN training_catalog tc ON tc.training_id=tr.training_id WHERE tr.requirement_id=?",(req_id,)).fetchone()
                if rr and et=="validation": label=f"Validation — {rr[0]}"
                elif rr and et=="action" and label in {"Action","Evidence Validation — Requirement"}: label=f"{label} — {rr[0]}"
            doc["linked_label"]=label; doc["linked_context"]=context
            mime=str(doc.get("mime_type") or "").lower()
            doc["viewable"]=1 if (mime.startswith("image/") or mime=="application/pdf") else 0

        timeline=[]
        for x in position_history:
            timeline.append({"event_date":x.get("effective_start") or x.get("changed_at") or "","event_type":"Position","title":x.get("position_name") or "Position Update","detail":x.get("change_reason") or x.get("change_type") or "Position record updated","actor":x.get("changed_by_name") or "System"})
        for x in training:
            actor_id=x.get("updated_by") or x.get("created_by") or ""; actor_row=con.execute("SELECT display_name FROM app_users WHERE user_id=?",(actor_id,)).fetchone() if actor_id else None
            timeline.append({"event_date":x.get("completion_date") or x.get("training_date") or "","event_type":"Training","title":x.get("training_name") or "Training","detail":x.get("result_status") or "Training record","actor":actor_row[0] if actor_row else (actor_id or "System")})
        for x in certifications:
            actor_id=x.get("updated_by") or x.get("created_by") or ""; actor_row=con.execute("SELECT display_name FROM app_users WHERE user_id=?",(actor_id,)).fetchone() if actor_id else None
            timeline.append({"event_date":x.get("issue_date") or "","event_type":"Certification","title":x.get("certification_name_raw") or "Certification","detail":x.get("certification_status") or "Certification record","actor":actor_row[0] if actor_row else (actor_id or "System")})
        run_rows = con.execute("""SELECT ar.run_id,ar.run_type,ar.completed_at,ar.covered_count,ar.potential_gap_count,ar.validation_required_count,ar.not_assessed_count,ar.notes,COALESCE(u.display_name,ar.initiated_by,'System') actor
                FROM assessment_runs ar LEFT JOIN app_users u ON u.user_id=ar.initiated_by WHERE ar.employee_pk=? ORDER BY ar.completed_at DESC""",(employee_pk,)).fetchall()
        for r in run_rows:
            detail=f"Covered {r[3] or 0} · Potential Gap {r[4] or 0} · Validation {r[5] or 0} · Not Assessed {r[6] or 0}"
            if r[7]: detail += f" · Trigger: {r[7]}"
            elif "reassess" in str(r[1] or "").lower(): detail += " · Trigger: employee marked Pending Reassessment after data/evidence change"
            timeline.append({"event_date":r[2] or "","event_type":"Assessment","title":r[1] or "Assessment","detail":detail,"actor":r[8] or "System"})
        action_events = con.execute("""SELECT ah.changed_at,ah.event_type,a.title,ah.old_status,ah.new_status,ah.remarks,COALESCE(u.display_name,ah.changed_by,'System') actor
                FROM action_history ah JOIN actions a ON a.action_id=ah.action_id LEFT JOIN app_users u ON u.user_id=ah.changed_by
                WHERE a.employee_pk=? AND COALESCE(a.is_deleted,0)=0 ORDER BY ah.changed_at DESC""",(employee_pk,)).fetchall()
        for r in action_events:
            # Hide meaningless no-op status rows; evidence upload is represented as its own event below.
            if (r[3] or "")== (r[4] or "") and not (r[5] or "") and str(r[1] or "").lower() in {"update","status update","updated"}:
                continue
            detail = f"{r[3] or ''} → {r[4] or ''}".strip(" →")
            if r[5]: detail += (" · " if detail else "") + str(r[5])
            timeline.append({"event_date":r[0] or "","event_type":"Action","title":r[2] or r[1] or "Action","detail":detail or r[1] or "Action updated","actor":r[6] or "System"})
        for doc in documents:
            detail=f"{doc.get('document_type') or 'Evidence'} · linked to {doc.get('linked_label') or doc.get('entity_type') or 'record'}"
            if doc.get("description"): detail += f" · {doc.get('description')}"
            timeline.append({"event_date":doc.get("uploaded_at") or "","event_type":"Evidence","title":f"Evidence Uploaded — {doc.get('original_filename') or 'file'}","detail":detail,"actor":doc.get("uploaded_by_name") or "System"})
        timeline.sort(key=lambda x: str(x.get("event_date") or ""), reverse=True)

    return {"employee":dict(emp),"summary":{"requirements":requirement_count,"training_history":training_count,"certifications":cert_count,
            "open_actions":open_actions,"overdue_actions":overdue,**{str(k).lower().replace(' ','_'):v for k,v in distribution.items()}},
            "assessment_distribution":distribution,"requirements":requirements,"training":training,"certifications":certifications,
            "competency":assessment_rows.to_dict(orient="records"),"actions":actions,"documents":documents,"position_history":position_history,"timeline":timeline[:200]}


@app.get("/api/employees/{employee_pk}/export")
def export_employee_profile(employee_pk: int):
    profile=employee_profile(employee_pk)
    emp=profile["employee"]
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        pd.DataFrame([{
            "Employee ID":emp.get("employee_code"),"Employee Name":emp.get("employee_name"),"Position":emp.get("position_name"),"Department":emp.get("department"),
            "Employment Status":emp.get("employment_status"),"Assessment State":emp.get("assessment_state"),"Join Date":emp.get("join_date"),"Last Assessment":emp.get("last_assessment_at")
        }]).to_excel(writer,index=False,sheet_name='OVERVIEW')
        pd.DataFrame(profile["requirements"]).to_excel(writer,index=False,sheet_name='REQUIREMENTS')
        pd.DataFrame(profile["training"]).to_excel(writer,index=False,sheet_name='TRAINING_HISTORY')
        pd.DataFrame(profile["certifications"]).to_excel(writer,index=False,sheet_name='CERTIFICATIONS')
        pd.DataFrame(profile["competency"]).to_excel(writer,index=False,sheet_name='ASSESSMENT')
        pd.DataFrame(profile["actions"]).to_excel(writer,index=False,sheet_name='ACTIONS')
        pd.DataFrame(profile["documents"]).to_excel(writer,index=False,sheet_name='DOCUMENTS')
        pd.DataFrame(profile["timeline"]).to_excel(writer,index=False,sheet_name='HISTORY')
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(max(12,max(len(str(c.value or '')) for c in col)+2),48)
    out.seek(0)
    user=current_user(); safe=''.join(ch if ch.isalnum() or ch in ('_','-') else '_' for ch in str(emp.get('employee_name') or employee_pk))
    filename=f"Employee_360_{safe}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    write_audit(user,'EXPORT_EMPLOYEE_360','Employee',str(employee_pk),filename)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}"'})


@app.get("/api/search")
def global_search(q: str = Query(min_length=2), limit: int = 20) -> list[dict[str, Any]]:
    qlike = f"%{q.strip()}%"; lim=max(1,min(limit,50)); results=[]
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory=sqlite3.Row
        rows=con.execute("""SELECT e.employee_pk,e.employee_name,e.employee_code,p.position_name FROM employees e LEFT JOIN positions p ON p.position_id=e.current_position_id
                 WHERE e.employee_name LIKE ? OR e.employee_code LIKE ? OR p.position_name LIKE ? ORDER BY e.active_flag DESC,e.employee_name LIMIT ?""",(qlike,qlike,qlike,lim)).fetchall()
        results += [{"type":"Employee","label":r["employee_name"],"subtitle":f"{r['employee_code'] or ''} · {r['position_name'] or 'No Position'}","employee_pk":r["employee_pk"],"entity_id":r["employee_pk"]} for r in rows]
        rows=con.execute("""SELECT c.certification_id,c.employee_pk,e.employee_name,c.certification_name_raw,COALESCE(c.certification_status,c.status) cert_status
                 FROM certifications c JOIN employees e ON e.employee_pk=c.employee_pk WHERE COALESCE(c.record_status,'Active')<>'Archived' AND (c.certification_name_raw LIKE ? OR e.employee_name LIKE ?) ORDER BY c.certification_id DESC LIMIT ?""",(qlike,qlike,lim)).fetchall()
        results += [{"type":"Certification","label":r["certification_name_raw"],"subtitle":f"{r['employee_name']} · {r['cert_status'] or ''}","employee_pk":r["employee_pk"],"entity_id":r["certification_id"]} for r in rows]
        rows=con.execute("""SELECT a.action_id,a.employee_pk,e.employee_name,a.title,a.status FROM actions a JOIN employees e ON e.employee_pk=a.employee_pk
                 WHERE COALESCE(a.is_deleted,0)=0 AND (a.title LIKE ? OR e.employee_name LIKE ?) ORDER BY a.action_id DESC LIMIT ?""",(qlike,qlike,lim)).fetchall()
        results += [{"type":"Action","label":r["title"],"subtitle":f"{r['employee_name']} · {r['status']}","employee_pk":r["employee_pk"],"entity_id":r["action_id"]} for r in rows]
        rows=con.execute("""SELECT training_id,training_name,category FROM training_catalog WHERE active_flag=1 AND training_name LIKE ? ORDER BY training_name LIMIT ?""",(qlike,lim)).fetchall()
        results += [{"type":"Training","label":r["training_name"],"subtitle":r["category"] or "Training Catalog","employee_pk":None,"entity_id":r["training_id"]} for r in rows]
    order={"Employee":1,"Certification":2,"Action":3,"Training":4}
    return sorted(results,key=lambda x:(order.get(x["type"],9),x["label"]))[:lim]


def _export_frame(dataset: str, search: str="", status: str="", employee_pk: int | None = None, position: str="", department: str="", requirement_type: str="", deleted: str="active") -> tuple[pd.DataFrame,str]:
    key=dataset.lower().strip().replace('-','_')
    if key=='employees':
        df=query_df("""SELECT e.employee_code AS [Employee ID],e.employee_name AS [Employee Name],p.position_name AS Position,COALESCE(e.department,p.department,'') AS Department,
                    e.employment_status AS [Employment Status],e.join_date AS [Join Date],e.exit_date AS [Exit Date],e.assessment_state AS [Assessment State],e.last_assessment_at AS [Last Assessment],e.remarks AS Remarks
                FROM employees e LEFT JOIN positions p ON p.position_id=e.current_position_id ORDER BY e.active_flag DESC,e.employee_name""")
        title='Employee'
    elif key=='training_requirements':
        employee_search_consumed=False
        if employee_pk is not None:
            employee_where='e.employee_pk=?'
            employee_params=(employee_pk,)
        else:
            employee_where=''
            employee_params=()
            if search:
                like=f"%{search.strip()}%"
                matches=query_df("""SELECT employee_pk FROM employees
                                  WHERE active_flag=1 AND (employee_name LIKE ? OR employee_code LIKE ?)
                                  ORDER BY employee_name""", (like,like))
                if not matches.empty:
                    employee_where='(e.employee_name LIKE ? OR e.employee_code LIKE ?)'
                    employee_params=(like,like)
                    employee_search_consumed=True
        if employee_where:
            # Employee-context export: one row per applicable requirement and employee.
            # Text search that resolves to employee name/ID uses the same semantics as
            # the Training Explorer screen instead of leaking other employees sharing
            # the same position.
            df=query_df(f"""SELECT e.employee_code AS [Employee ID],e.employee_name AS Employee,
                        p.position_name AS Position,COALESCE(e.department,p.department,'') AS Department,
                        tr.requirement_id AS [Requirement ID],tc.training_name AS Training,
                        tr.requirement_type AS [Requirement Type],tr.requirement_code AS Code,
                        tr.delivery_type AS Delivery,tr.regulatory_flag AS Regulatory,
                        tr.requirement_status AS Status,tr.effective_from AS [Effective From],
                        tr.effective_to AS [Effective To],tr.requirement_source AS Source,tr.remarks AS Remarks
                    FROM employees e
                    LEFT JOIN positions p ON p.position_id=e.current_position_id
                    JOIN training_requirements tr ON COALESCE(tr.position_standard_id,tr.position_id)=e.current_position_id
                    JOIN training_catalog tc ON tc.training_id=tr.training_id
                    WHERE {employee_where} AND e.active_flag=1
                    ORDER BY e.employee_name,tc.training_name""", employee_params)
            title='Training Requirement — Employee Context'
            if employee_search_consumed:
                search=''
        else:
            df=query_df("""SELECT tr.requirement_id AS [Requirement ID],p.position_name AS Position,
                        COALESCE(p.department,'') AS Department,tc.training_name AS Training,
                        tr.requirement_type AS [Requirement Type],tr.requirement_code AS Code,
                        tr.delivery_type AS Delivery,tr.regulatory_flag AS Regulatory,
                        tr.requirement_status AS Status,tr.effective_from AS [Effective From],
                        tr.effective_to AS [Effective To],tr.requirement_source AS Source,tr.remarks AS Remarks
                    FROM training_requirements tr
                    LEFT JOIN positions p ON p.position_id=COALESCE(tr.position_standard_id,tr.position_id)
                    LEFT JOIN training_catalog tc ON tc.training_id=tr.training_id
                    ORDER BY p.position_name,tc.training_name""")
            title='Training Requirement'
        if position and not df.empty and 'Position' in df.columns:
            df=df[df['Position'].astype(str).str.lower()==position.lower()]
        if department and not df.empty and 'Department' in df.columns:
            df=df[df['Department'].astype(str).str.lower()==department.lower()]
        if requirement_type and not df.empty and 'Requirement Type' in df.columns:
            df=df[df['Requirement Type'].astype(str).str.lower()==requirement_type.lower()]
    elif key=='training_history':
        df=query_df("""SELECT e.employee_code AS [Employee ID],e.employee_name AS Employee,p.position_name AS Position,COALESCE(th.training_name,th.training_name_raw,tc.training_name) AS Training,
                    th.training_date AS [Training Date],th.completion_date AS [Completion Date],COALESCE(th.result_status,th.completion_status) AS Result,th.provider AS Provider,th.valid_until AS [Valid Until],th.certificate_reference AS [Certificate Reference],th.evidence_status AS Evidence,th.remarks AS Remarks
                FROM training_history th JOIN employees e ON e.employee_pk=th.employee_pk LEFT JOIN positions p ON p.position_id=e.current_position_id LEFT JOIN training_catalog tc ON tc.training_id=COALESCE(th.training_id,th.training_pk)
                WHERE COALESCE(th.record_status,'Active')<>'Archived' ORDER BY COALESCE(th.completion_date,th.training_date) DESC""")
        title='Training History'
    elif key=='certifications':
        df=query_df("""SELECT e.employee_code AS [Employee ID],e.employee_name AS Employee,p.position_name AS Position,c.certification_name_raw AS Certification,c.certificate_number AS [Certificate Number],COALESCE(c.issuer,c.issuer_type) AS Issuer,
                    COALESCE(c.issue_date,c.certification_date) AS [Issue Date],COALESCE(c.expiry_date,c.expired_date) AS [Expiry Date],
                    CASE WHEN lower(COALESCE(c.certification_status,c.status,'')) LIKE '%proses%' THEN 'In Process'
                         WHEN COALESCE(c.expiry_date,c.expired_date) IS NOT NULL AND date(COALESCE(c.expiry_date,c.expired_date))<date('now') THEN 'Expired'
                         WHEN COALESCE(c.expiry_date,c.expired_date) IS NOT NULL AND date(COALESCE(c.expiry_date,c.expired_date))<=date('now','+90 day') THEN 'Near Expiry'
                         ELSE 'Active' END AS [Certification Status],c.renewal_status AS [Renewal Status],c.reason_code AS Reason,c.remarks AS Remarks
                FROM certifications c JOIN employees e ON e.employee_pk=c.employee_pk LEFT JOIN positions p ON p.position_id=e.current_position_id WHERE COALESCE(c.record_status,'Active')<>'Archived' ORDER BY COALESCE(c.expiry_date,c.expired_date,'9999-12-31')""")
        title='Certification'
    elif key=='competency':
        df=query_df("""SELECT da.employee_name AS Employee,da.position_name AS Position,da.training_name AS Requirement,da.requirement_type AS [Requirement Type],da.mapping_category AS Mapping,
                    da.final_status AS [Coverage Status],da.priority AS Priority,da.assessment_source AS [Assessment Source],da.evidence_summary AS Evidence,da.rationale AS Rationale,da.decision_action AS Recommendation,da.evaluated_at AS [Evaluated At]
                FROM decision_assessment_current da ORDER BY da.employee_name,da.training_name""")
        if df.empty:
            df=query_df("""SELECT ca.employee_pk,ca.training_name AS Requirement,ca.position_tna AS Position,ca.coverage_status AS [Coverage Status],ca.priority AS Priority,ca.recommendation AS Recommendation,ca.assessment_source AS [Assessment Source] FROM competency_assessment ca""")
        title='Competency Assessment'
    elif key=='actions':
        action_scope={"active":"COALESCE(a.is_deleted,0)=0","deleted":"COALESCE(a.is_deleted,0)=1","all":"1=1"}.get(deleted)
        if not action_scope: raise HTTPException(status_code=400,detail="deleted harus active, deleted, atau all.")
        df=query_df(f"""SELECT a.action_id AS [Action ID],e.employee_name AS Employee,p.position_name AS Position,tc.training_name AS Requirement,a.action_type AS [Action Type],a.title AS Action,a.priority AS Priority,u.display_name AS PIC,
                    a.due_date AS [Due Date],a.status AS Status,CASE WHEN a.is_deleted=1 THEN 'Deleted' ELSE 'Active' END AS [Record State],a.deletion_reason AS [Deletion Reason],a.deleted_at AS [Deleted At],a.deleted_by AS [Deleted By],a.reason_code AS Reason,a.remarks AS Remarks,a.created_at AS [Created At],a.completed_at AS [Completed At],a.closed_at AS [Closed At]
                FROM actions a JOIN employees e ON e.employee_pk=a.employee_pk LEFT JOIN positions p ON p.position_id=e.current_position_id LEFT JOIN training_requirements tr ON tr.requirement_id=a.requirement_id LEFT JOIN training_catalog tc ON tc.training_id=tr.training_id LEFT JOIN app_users u ON u.user_id=a.pic_user_id WHERE {action_scope} ORDER BY a.action_id DESC""")
        title='Action'
    else:
        raise HTTPException(status_code=404,detail="Dataset export tidak dikenal.")
    if search and not df.empty:
        mask=df.astype(str).apply(lambda c:c.str.contains(search,case=False,regex=False)).any(axis=1); df=df[mask]
    if status and not df.empty:
        cols=[c for c in df.columns if 'status' in c.lower() or c.lower()=='priority']
        if cols:
            mask=df[cols].astype(str).apply(lambda c:c.str.lower().eq(status.lower())).any(axis=1); df=df[mask]
    return df.fillna(""),title


@app.get("/api/export/{dataset}")
def export_dataset(dataset: str, search: str="", status: str="", employee_pk: int | None = None, position: str="", department: str="", requirement_type: str="", deleted: str="active"):
    user=current_user(); df,title=_export_frame(dataset,search,status,employee_pk,position,department,requirement_type,deleted)
    employee_label="All"
    if employee_pk is not None:
        with sqlite3.connect(DB_PATH) as con:
            row=con.execute("SELECT employee_name FROM employees WHERE employee_pk=?",(employee_pk,)).fetchone()
            employee_label=row[0] if row else f"Employee #{employee_pk}"
    meta=pd.DataFrame([{"Exported At":datetime.now().isoformat(timespec='seconds'),"Dataset":title,"Search Filter":search or "All","Status Filter":status or "All","Employee Filter":employee_label,"Position Filter":position or "All","Department Filter":department or "All","Requirement Type Filter":requirement_type or "All","Deleted Filter":deleted,"Records":len(df),"Build":BUILD}])
    defs=pd.DataFrame([
        {"Term":"Potential Gap","Definition":"Requirement berlaku dan evidence pemenuhan yang diharapkan belum ditemukan; bukan competency gap final."},
        {"Term":"Validation Required","Definition":"Hubungan mapping/evidence belum cukup pasti dan perlu validasi TCD."},
        {"Term":"Not Assessed","Definition":"Data atau mapping belum memadai untuk assessment."},
        {"Term":"Covered","Definition":"Evidence yang relevan dan valid mendukung pemenuhan requirement."},
        {"Term":"Pending Initial Assessment","Definition":"Data awal employee baru belum dikonfirmasi atau assessment awal belum dijalankan."},
        {"Term":"Pending Reassessment","Definition":"Evidence/position/requirement berubah setelah assessment dan perlu dihitung ulang."},
    ])
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        df.to_excel(writer,index=False,sheet_name='DATA')
        meta.to_excel(writer,index=False,sheet_name='EXPORT_INFO')
        defs.to_excel(writer,index=False,sheet_name='STATUS_DEFINITION')
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            for col in ws.columns:
                width=min(max(12,max(len(str(c.value or '')) for c in col)+2),48); ws.column_dimensions[col[0].column_letter].width=width
    out.seek(0); stamp=datetime.now().strftime('%Y%m%d_%H%M%S'); filename=f"{dataset}_{stamp}.xlsx"
    write_audit(user,"EXPORT_DATA",title,filename,f"records={len(df)}; search={search or 'All'}; status={status or 'All'}; employee_pk={employee_pk or 'All'}; position={position or 'All'}; department={department or 'All'}; requirement_type={requirement_type or 'All'}")
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}"'})

# -----------------------------------------------------------------------------
# Build v22.10.5 — Reporting & Management Analytics retained
# -----------------------------------------------------------------------------

@app.get("/api/analytics/departments")
def analytics_departments() -> list[str]:
    with sqlite3.connect(DB_PATH) as con:
        return list_departments(con)


@app.get("/api/analytics/management")
def analytics_management(department: str = "") -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        data = management_analytics(con, department=department or None, refresh_snapshot=not bool(department), created_by=user["user_id"])
        con.commit()
    return data


@app.get("/api/analytics/trends")
def analytics_trends(months: int = 12) -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        return trend_analytics(con, months=months)


@app.post("/api/analytics/snapshot")
def analytics_snapshot() -> dict[str, Any]:
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        result = refresh_monthly_snapshot(con, created_by=user["user_id"])
        con.commit()
    write_audit(user, "REFRESH_ANALYTICS_SNAPSHOT", "Reporting", result.get("period_key", ""), "Monthly current-period snapshot refreshed")
    return {"status": "success", "snapshot": result}


def _style_management_workbook(writer) -> None:
    # The application already uses openpyxl for Excel delivery. Apply a restrained management-report style.
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    header_fill = PatternFill("solid", fgColor="173A56")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="EAF1F6")
    thin = Side(style="thin", color="D8E2EA")
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in ws.columns:
            values = [len(str(c.value or "")) for c in col]
            width = min(max(12, max(values or [12]) + 2), 42)
            ws.column_dimensions[col[0].column_letter].width = width
        if ws.title in {"EXECUTIVE_SUMMARY", "REPORT_INFO", "DEFINITIONS"} and ws.max_row >= 2:
            for cell in ws[2]:
                if cell.value is not None:
                    cell.fill = sub_fill


@app.get("/api/reports/management/export")
def export_management_report(department: str = ""):
    user = current_user()
    with sqlite3.connect(DB_PATH) as con:
        data = management_analytics(con, department=department or None, refresh_snapshot=not bool(department), created_by=user["user_id"])
        con.commit()
    frames = analytics_frames(data)
    info = pd.DataFrame([{
        "Generated At": data["generated_at"], "Department": data["department"], "Build": BUILD,
        "Report Scope": "Current operational state + available monthly snapshots",
        "Important Note": "Unavailable rates remain blank when the required historical/Training History data does not yet exist."
    }])
    readiness_rows=[]
    for key,val in data.get("data_readiness",{}).items():
        readiness_rows.append({"Area":key,"Ready":bool(val.get("ready")),"Records/Points":val.get("records",val.get("points","")),"Note":val.get("note","")})
    frames["DATA_READINESS"] = pd.DataFrame(readiness_rows)
    frames["REPORT_INFO"] = info
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for sheet_name, df in frames.items():
            frame = df if not df.empty else pd.DataFrame([{"Information":"No records available for this section."}])
            frame.to_excel(writer, index=False, sheet_name=sheet_name[:31])
        _style_management_workbook(writer)
    out.seek(0)
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
    dep_slug=(department or 'All').strip().replace(' ','_').replace('/','-')
    filename=f"Management_Analytics_{dep_slug}_{stamp}.xlsx"
    with sqlite3.connect(DB_PATH) as con:
        con.execute("INSERT INTO management_report_log(period_label,department,report_format,generated_by,record_summary,filename) VALUES (?,?,?,?,?,?)",
                    (date.today().strftime('%Y-%m'), department or 'All', 'xlsx', user['user_id'], json.dumps(data['summary'],ensure_ascii=False), filename))
        con.commit()
    write_audit(user, "EXPORT_MANAGEMENT_REPORT", "Reporting", filename, f"department={department or 'All'}")
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition":f'attachment; filename="{filename}"'})


@app.get("/api/dashboard")
def dashboard() -> dict[str, Any]:
    with sqlite3.connect(DB_PATH) as con:
        employee_count=int(con.execute("SELECT COUNT(*) FROM employees WHERE active_flag=1").fetchone()[0])
        cert_count=int(con.execute("SELECT COUNT(*) FROM certifications WHERE COALESCE(record_status,'Active')<>'Archived'").fetchone()[0])
        history_count=int(con.execute("SELECT COUNT(*) FROM training_history WHERE COALESCE(record_status,'Active')<>'Archived'").fetchone()[0])
        requirement_count=int(con.execute("SELECT COUNT(*) FROM training_requirements WHERE active_flag=1 AND COALESCE(requirement_status,'Active')='Active'").fetchone()[0])
        expired_count=int(con.execute("SELECT COUNT(*) FROM certifications WHERE COALESCE(record_status,'Active')<>'Archived' AND COALESCE(expiry_date,expired_date) IS NOT NULL AND date(COALESCE(expiry_date,expired_date))<date('now')").fetchone()[0])
        near_expiry_count=int(con.execute("SELECT COUNT(*) FROM certifications WHERE COALESCE(record_status,'Active')<>'Archived' AND COALESCE(expiry_date,expired_date) IS NOT NULL AND date(COALESCE(expiry_date,expired_date))>=date('now') AND date(COALESCE(expiry_date,expired_date))<=date('now','+90 day')").fetchone()[0])
        urgent_expiry=int(con.execute("SELECT COUNT(*) FROM certifications WHERE COALESCE(record_status,'Active')<>'Archived' AND COALESCE(expiry_date,expired_date) IS NOT NULL AND date(COALESCE(expiry_date,expired_date))>=date('now') AND date(COALESCE(expiry_date,expired_date))<=date('now','+30 day')").fetchone()[0])
        status=pd.read_sql_query("SELECT coverage_status AS status,COUNT(*) total FROM v_coverage_gap_results GROUP BY coverage_status ORDER BY total DESC",con).fillna("").to_dict(orient='records')
        priority=pd.read_sql_query("SELECT priority,COUNT(*) total FROM v_coverage_gap_results GROUP BY priority ORDER BY total DESC",con).fillna("").to_dict(orient='records')
        lifecycle=pd.read_sql_query("SELECT COALESCE(assessment_state,'Assessed') state,COUNT(*) total FROM employees WHERE active_flag=1 GROUP BY COALESCE(assessment_state,'Assessed')",con).fillna("").to_dict(orient='records')
        lm={str(x['state']):int(x['total']) for x in lifecycle}
        open_actions=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND status NOT IN ('Completed','Closed','Cancelled','No Action Required')").fetchone()[0])
        overdue_actions=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND status NOT IN ('Completed','Closed','Cancelled','No Action Required') AND due_date IS NOT NULL AND date(due_date)<date('now')").fetchone()[0])
        high_actions=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND priority='High' AND status NOT IN ('Completed','Closed','Cancelled','No Action Required')").fetchone()[0])
        unassigned=int(con.execute("SELECT COUNT(*) FROM actions WHERE COALESCE(is_deleted,0)=0 AND pic_user_id IS NULL AND status NOT IN ('Completed','Closed','Cancelled','No Action Required')").fetchone()[0])
        potential_high=int(con.execute("SELECT COUNT(*) FROM v_coverage_gap_results WHERE coverage_status='Potential Gap' AND priority='High'").fetchone()[0])
        validation_count=int(con.execute("SELECT COUNT(*) FROM v_coverage_gap_results WHERE coverage_status='Validation Required'").fetchone()[0])
        unread=int(con.execute("SELECT COUNT(*) FROM notifications WHERE active_flag=1 AND read_flag=0 AND dismissed_at IS NULL").fetchone()[0])
        action_progress=pd.read_sql_query("SELECT status,COUNT(*) total FROM actions WHERE COALESCE(is_deleted,0)=0 GROUP BY status ORDER BY total DESC",con).fillna("").to_dict(orient='records')
        upcoming=pd.read_sql_query("""SELECT c.certification_id,c.employee_pk,e.employee_name,c.certification_name_raw certification_name,COALESCE(c.expiry_date,c.expired_date) expiry_date,
                       CAST(julianday(date(COALESCE(c.expiry_date,c.expired_date)))-julianday(date('now')) AS INTEGER) days_remaining,COALESCE(c.renewal_status,'') renewal_status
                    FROM certifications c JOIN employees e ON e.employee_pk=c.employee_pk
                    WHERE e.active_flag=1 AND COALESCE(c.record_status,'Active')<>'Archived' AND COALESCE(c.expiry_date,c.expired_date) IS NOT NULL
                      AND date(COALESCE(c.expiry_date,c.expired_date))>=date('now') AND date(COALESCE(c.expiry_date,c.expired_date))<=date('now','+180 day')
                    ORDER BY date(COALESCE(c.expiry_date,c.expired_date)) LIMIT 8""",con).fillna("").to_dict(orient='records')
        for item in upcoming:
            validity='Near Expiry' if int(item.get('days_remaining') or 9999) <= 90 else 'Active'
            item['renewal_status']=renewal_status_for(validity,item.get('renewal_status'))
        recent=pd.read_sql_query("""SELECT al.created_at,al.action,al.dataset_type,al.target,al.details,COALESCE(u.display_name,al.user_id) user_name
                     FROM audit_log al LEFT JOIN app_users u ON u.user_id=al.user_id ORDER BY al.audit_id DESC LIMIT 10""",con).fillna("").to_dict(orient='records')
        needs=[
            {"key":"expired","label":"Expired Certifications","count":expired_count,"severity":"critical","page":"certification","filter":"Expired"},
            {"key":"potential_high","label":"High Priority Potential Gap","count":potential_high,"severity":"high","page":"competency","filter":"Potential Gap"},
            {"key":"overdue","label":"Overdue Actions","count":overdue_actions,"severity":"critical","page":"actions","filter":"overdue"},
            {"key":"urgent_expiry","label":"Certification ≤30 Days","count":urgent_expiry,"severity":"high","page":"certification","filter":"Near Expiry"},
            {"key":"validation","label":"Awaiting Validation","count":validation_count,"severity":"attention","page":"validation","filter":""},
            {"key":"pending_reassessment","label":"Pending Reassessment","count":lm.get('Pending Reassessment',0),"severity":"attention","page":"assessment","filter":""},
            {"key":"unassigned","label":"Unassigned Actions","count":unassigned,"severity":"info","page":"actions","filter":"unassigned"},
        ]
    return {"kpi":{"employees":employee_count,"active_requirements":requirement_count,"certifications":cert_count,"expired":expired_count,"near_expiry_90d":near_expiry_count,
                    "urgent_expiry_30d":urgent_expiry,"training_history":history_count,"pending_initial":lm.get('Pending Initial Assessment',0),"assessment_ready":lm.get('Assessment Ready',0),
                    "pending_reassessment":lm.get('Pending Reassessment',0),"open_actions":open_actions,"high_actions":high_actions,"overdue_actions":overdue_actions,"unread_notifications":unread},
            "assessment_lifecycle":lifecycle,"decision_status":status,"priority":priority,"needs_attention":needs,"action_progress":action_progress,"upcoming_certifications":upcoming,"recent_activity":recent}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("api_server:app", host="0.0.0.0", port=8000, reload=False)
